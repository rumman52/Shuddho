"""CI-only native engine QA. Requires LibreOffice and pdftoppm; no model calls.

Run with PYTHONPATH=.:tests python tests/verify_office_native.py /tmp/office-qa
The production image does not require an Office process or this fixture.
"""
import json
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from coworker_samples import work_draft
from services.coworker.exports import render_in_subprocess
from services.coworker.schemas import parse_draft


def main(folder):
    office = shutil.which("libreoffice") or shutil.which("soffice")
    if not office or not shutil.which("pdftoppm"):
        raise RuntimeError("Native QA requires LibreOffice and pdftoppm; this gate cannot be skipped")
    folder.mkdir(parents=True, exist_ok=True)
    originals, rendered, edited, recalculated = [folder / name for name in ("originals", "rendered", "edited", "recalculated")]
    for path in (originals, rendered, edited, recalculated):
        path.mkdir(exist_ok=True)
    profile = folder / "office-profile"

    def convert(files, extension, destination):
        result = subprocess.run([office, "--headless", "-env:UserInstallation=" + profile.as_uri(),
            "--convert-to", extension, "--outdir", str(destination), *map(str, files)],
            capture_output=True, text=True, timeout=60)
        if result.returncode:
            raise RuntimeError(result.stderr[-1000:])
        outputs = [destination / (file.stem + "." + extension) for file in files]
        if any(not path.is_file() for path in outputs):
            raise RuntimeError("Office did not produce every requested file: " + result.stdout[-1000:])
        return outputs

    expected_slides = {}
    for language in ("en", "bn", "ar"):
        for skill, suffix in (("presentation", "pptx"), ("spreadsheet", "xlsx")):
            draft = work_draft(skill, language)
            files = dict((name, body) for name, _, body in render_in_subprocess(draft, [], skill_id=skill))
            path = originals / f"{skill}-{language}.{suffix}"
            path.write_bytes(files[f"{skill}.{suffix}"])
            (originals / f"{skill}-{language}-handout.pdf").write_bytes(files[f"{skill}.pdf"])
            if skill == "presentation": expected_slides[path.stem] = len(draft.slides)
    # Exercise the actual text limits with wide CJK characters, including charts.
    value = work_draft("presentation").model_dump()
    value["output_language"] = "zh"
    value["title"] = "跨语言项目报告"
    for slide in value["slides"]:
        slide["title"] = ("项目进度与后续工作计划" * 6)[:60]
        slide["bullets"] = [("团队已经完成审查并记录结果请根据资料安排后续工作" * 6)[:110]] * (3 if slide["layout"] == "content" else 1)
    draft = parse_draft("presentation", value)
    path = originals / "presentation-zh-dense.pptx"
    path.write_bytes(next(body for name, _, body in render_in_subprocess(draft, [], skill_id="presentation") if name.endswith(".pptx")))
    expected_slides[path.stem] = len(draft.slides)
    value = work_draft("spreadsheet").model_dump()
    value.update(output_language="zh", title="跨语言项目报告" * 15, summary="根据提供的资料核对数据并记录后续计划。" * 25,
                 columns=[{"id": "text", "label": "项目说明与后续工作" * 5, "format": "text"}],
                 rows=[["团队已经完成审查并记录结果请根据资料安排后续工作" * 6]], calculations=[], chart=None, summary_label="")
    draft = parse_draft("spreadsheet", value)
    (originals / "spreadsheet-zh-long.xlsx").write_bytes(next(body for name, _, body in render_in_subprocess(draft, [], skill_id="spreadsheet") if name.endswith(".xlsx")))
    pdfs = convert(sorted(originals.glob("*.pptx")) + sorted(originals.glob("*.xlsx")), "pdf", rendered)
    for pdf in pdfs:
        reader = PdfReader(pdf)
        assert len(reader.pages) == expected_slides.get(pdf.stem, len(reader.pages))
        assert all(page.extract_text().strip() for page in reader.pages), pdf.name
        if pdf.stem == "presentation-ar":
            assert "02 / 03" in reader.pages[1].extract_text()
        if pdf.stem == "spreadsheet-en":
            assert len(reader.pages) == 2
            # Chart categories and the highest tick must be on the same page.
            chart_text = reader.pages[1].extract_text()
            assert all(value in chart_text for value in ("A", "B", "C", "40"))
        subprocess.run(["pdftoppm", "-scale-to", "1400", "-png", str(pdf), str(rendered / pdf.stem)], check=True, capture_output=True, timeout=30)
    outcomes = {}
    cases = {"changed": (7, 87.5, 103.5), "zero": (0, 0, 16), "missing": (None, None, None), "negative": (-3, -37.5, -21.5)}
    for name, (quantity, _, _) in cases.items():
        book = load_workbook(originals / "spreadsheet-en.xlsx")
        book.active["B6"] = quantity
        book.save(edited / (name + ".xlsx"))
        book.close()
    converted = convert(sorted(edited.glob("*.xlsx")), "xlsx", recalculated)
    for path in converted:
        _, expected_cost, expected_total = cases[path.stem]
        book = load_workbook(path, data_only=True)
        sheet = book.active
        blank = lambda value: None if value == "" else value
        assert blank(sheet["D6"].value) == expected_cost, (path.name, sheet["D6"].value)
        assert blank(sheet["D11"].value) == expected_total, (path.name, sheet["D11"].value)
        assert blank(sheet["E8"].value) is None
        assert not any(cell.data_type == "e" for row in sheet for cell in row)
        book.close()
        formulas = load_workbook(path, data_only=False)
        assert formulas.active["D6"].data_type == "f"
        assert len(formulas.active._charts) == 1
        chart = formulas.active._charts[0].series[0].val.numRef
        assert "D6" in chart.f.replace("$", "") and "D8" in chart.f.replace("$", "")
        if expected_cost is not None:
            assert float(chart.numCache.pt[0].v) == expected_cost
        formulas.close()
        outcomes[path.stem] = {"cost": expected_cost, "total": expected_total, "recalculated": True}
    (folder / "native-results.json").write_text(json.dumps({"rendered": [path.name for path in pdfs], "formula_edits": outcomes}, indent=2))
    shutil.rmtree(profile, ignore_errors=True)
    print("Native Office QA passed: 4 PPTX decks and 4 XLSX workbooks rendered; changed, zero, missing, and negative inputs recalculated; formula cells and live chart references preserved.")


if __name__ == "__main__":
    main(Path(sys.argv[1]).resolve())
