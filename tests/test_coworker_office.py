"""Editable outputs, arithmetic invariants, intake limits, and rollout recovery."""
import asyncio
import csv
import io
import json
import zipfile
from dataclasses import replace

import pytest

pytest.importorskip("sqlalchemy")
pytest.importorskip("pptx")
pytest.importorskip("openpyxl")

from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pypdf import PdfReader

from test_coworker import container, signed_client, account
from coworker_samples import work_draft, WorkModel
from services.coworker.calculations import table_values
from services.coworker.errors import CoworkerError
from services.coworker.exports import render_in_subprocess
from services.coworker.extraction import extract_in_subprocess, extract_text
from services.coworker.runner import DocumentRunner
from services.coworker.schemas import TaskCreate, UploadRequest, parse_draft


def enabled(container, value=True):
    container.settings = replace(container.settings, artifact_services_enabled=value)
    container.repository.settings = container.settings


@pytest.mark.parametrize("skill", ["presentation", "spreadsheet"])
def test_artifact_gate_is_independent_and_checkpoint_survives_disable(container, signed_client, skill):
    client, headers = signed_client
    payload = TaskCreate(skill_id=skill, instruction="Create from these supplied details", notes="A: 3 at 12.5; B: 2 at 8; C: 0 at 15.")
    owner = account(container)
    with pytest.raises(CoworkerError):
        container.repository.create_task(owner, payload, "office-request-key")
    with pytest.raises(CoworkerError):
        container.repository.create_upload(owner, UploadRequest(filename="source.xlsx", byte_size=100, sha256="a" * 64))
    assert client.get("/api/v1/me", headers=headers()).json()["usage"]["tasks_today"] == 0
    enabled(container)
    catalog = client.get("/api/v1/skills", headers=headers()).json()
    assert [item["id"] for item in catalog["skills"]] == ["report_email", "presentation", "spreadsheet"]
    assert catalog["upload_formats"] == ["txt", "docx", "pdf", "csv", "xlsx", "pptx"]
    task, _ = container.repository.create_task(owner, payload, "office-request-key")
    model = WorkModel()
    runner = DocumentRunner(container, model)
    asyncio.run(runner.phase(task["id"], "extract"))
    asyncio.run(runner.phase(task["id"], "draft"))
    enabled(container, False)
    asyncio.run(DocumentRunner(container, model).run_for_test(task["id"]))
    asyncio.run(DocumentRunner(container, model).run_for_test(task["id"]))
    saved = container.repository.get_task(owner, task["id"])
    assert model.calls == 1 and saved["state"] == "completed" and len(saved["artifacts"]) == 4
    assert bool(saved["preview"]) == (skill == "spreadsheet")
    if skill == "spreadsheet":
        assert saved["preview"]["rows"][0][3] == 37.5
    assert container.repository.create_task(owner, payload, "office-request-key")[0]["id"] == task["id"]
    with pytest.raises(CoworkerError):
        container.repository.create_task(owner, payload, "another-office-key")
    for artifact in saved["artifacts"]:
        with pytest.raises(CoworkerError) as error:
            container.repository.artifact(account(container, "bob"), artifact["id"])
        assert error.value.status_code == 404


@pytest.mark.parametrize("language", ["en", "bn", "ar"])
@pytest.mark.parametrize("skill", ["presentation", "spreadsheet"])
def test_editable_files_and_intake_roundtrip(skill, language):
    draft = work_draft(skill, language)
    files = {name: body for name, _, body in render_in_subprocess(draft, [{"id": "notes", "label": "Provided notes"}], skill_id=skill)}
    assert len(PdfReader(io.BytesIO(files[skill + ".pdf"])).pages) > 0
    if skill == "presentation":
        deck = Presentation(io.BytesIO(files["presentation.pptx"]))
        assert len(deck.slides) == 3
        assert draft.slides[1].title in [shape.text for shape in deck.slides[1].shapes if shape.has_text_frame]
        assert deck.slides[1].notes_slide.notes_text_frame.text == draft.slides[1].speaker_notes
        assert deck.slides[1].shapes[0].text_frame.paragraphs[0]._p.get_or_add_pPr().get("rtl") == "0"
        chart = next(shape.chart for shape in deck.slides[2].shapes if shape.has_chart)
        assert list(chart.series[0].values) == [12, 8]
        for slide in deck.slides:
            for shape in slide.shapes:
                assert shape.left >= 0 and shape.top >= 0
                assert shape.left + shape.width <= deck.slide_width and shape.top + shape.height <= deck.slide_height
        assert "saved chart data" in extract_in_subprocess(files["presentation.pptx"], "pptx", 20000)
    else:
        book = load_workbook(io.BytesIO(files["spreadsheet.xlsx"]), data_only=False)
        sheet = book.active
        assert sheet["D6"].data_type == "f" and "B6*C6" in sheet["D6"].value
        assert sheet["A6"].value == "A" and sheet["B6"].value == 3
        assert sheet.freeze_panes == "A6" and len(sheet._charts) == 1
        assert "A2:E2" in {str(merge) for merge in sheet.merged_cells}
        assert not sheet.sheet_properties.pageSetUpPr.fitToPage and sheet.row_breaks.brk[0].id == 12
        assert sheet["A6"].alignment.horizontal == ("right" if language == "ar" else "left")
        assert sheet.sheet_view.rightToLeft == (True if language == "ar" else None)
        cached = load_workbook(io.BytesIO(files["spreadsheet.xlsx"]), data_only=True).active
        assert cached["D6"].value == 37.5 and cached["D7"].value == 16 and cached["D8"].value == 0
        assert cached["E8"].value is None  # 15 / 0 must not become zero.
        assert cached["D11"].value == 53.5
        snapshot = extract_in_subprocess(files["spreadsheet.xlsx"], "xlsx", 20000)
        assert "37.5" in snapshot and "not recalculated" in snapshot
        rows = list(csv.reader(io.StringIO(files["spreadsheet.csv"].decode("utf-8-sig"))))
        assert rows[1][3] == "37.5" and rows[3][4] == ""


def test_missing_negative_and_undefined_inputs_remain_distinct():
    draft = work_draft("spreadsheet").model_dump()
    draft["rows"] = [["Refund", -2, 12.5], ["Unknown", None, 8], ["Zero", 0, 15]]
    draft["chart"] = None
    table = table_values(parse_draft("spreadsheet", draft))
    assert [row[3] for row in table["rows"]] == [-25, None, 0]
    assert [row[4] for row in table["rows"]] == [-6.25, None, None]
    assert table["totals"][1] is None and table["totals"][3] is None


@pytest.mark.parametrize("change", ["cycle", "formula", "wrong_type", "large", "duplicate", "width", "row_limit", "nan", "sheet_name"])
def test_model_cannot_bypass_spreadsheet_contract(change):
    value = work_draft("spreadsheet").model_dump()
    if change == "cycle": value["calculations"][0]["inputs"] = ["ratio", "quantity"]
    if change == "formula": value["calculations"][0]["operation"] = "WEBSERVICE"
    if change == "wrong_type": value["rows"][0][1] = "=SUM(1,2)"
    if change == "large": value["rows"][0][1:3] = [1e12, 1e12]
    if change == "duplicate": value["columns"][1]["id"] = "item"
    if change == "width": value["rows"][0].append(4)
    if change == "row_limit": value["rows"] *= 14
    if change == "nan": value["rows"][0][2] = float("nan")
    if change == "sheet_name": value["sheet_name"] = "invalid/name"
    with pytest.raises(ValueError): parse_draft("spreadsheet", value)


def test_text_injection_stays_text_in_native_file_and_csv():
    value = work_draft("spreadsheet").model_dump()
    value["rows"][0][0] = '=HYPERLINK("https://example.test", "click")'
    value["rows"][1][0] = "  @SUM(1,2)"
    value["columns"][0]["label"] = "+danger"
    value["chart"] = None
    files = {n: b for n, _, b in render_in_subprocess(parse_draft("spreadsheet", value), [], skill_id="spreadsheet")}
    sheet = load_workbook(io.BytesIO(files["spreadsheet.xlsx"])).active
    assert sheet["A6"].data_type == "s" and sheet["A6"].hyperlink is None
    assert sheet["A5"].data_type == "s"
    rows = list(csv.reader(io.StringIO(files["spreadsheet.csv"].decode("utf-8-sig"))))
    assert rows[0][0] == "'+danger" and rows[1][0].startswith("'=") and rows[2][0].startswith("'@")


def workbook_bytes(book):
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_uploads_reject_missing_formula_caches_oversized_coordinates_and_unsafe_archives():
    book = Workbook()
    book.active["A1"] = "=2+2"
    with pytest.raises(CoworkerError, match="no saved result"):
        extract_text(workbook_bytes(book), "xlsx")
    book.active["A1"] = "Input"
    book.active["A1000000"] = 4
    with pytest.raises(CoworkerError, match="200 data rows"):
        extract_text(workbook_bytes(book), "xlsx")
    for name, content in [("xl/vbaProject.bin", b"not executable"), ("../escape.xml", b"<xml/>"), ("unsafe.xml", b'<!DOCTYPE foo [<!ENTITY x "test">]><foo/>')]:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
            archive.writestr(name, content)
        with pytest.raises(CoworkerError): extract_text(buffer.getvalue(), "xlsx")
    with pytest.raises(CoworkerError, match="200 data rows"):
        extract_text(("a,b\n" * 202).encode(), "csv")
    with pytest.raises(CoworkerError, match="same number"):
        extract_text(b"a,b\n1,2,3", "csv")


def test_csv_preserves_multilingual_quoted_values():
    text = extract_in_subprocess('নাম,পরিমাণ\n"কাগজ, কলম",12\n'.encode(), "csv", 20000)
    assert json.loads(text.split("\n", 1)[1]) == [["নাম", "পরিমাণ"], ["কাগজ, কলম", "12"]]
