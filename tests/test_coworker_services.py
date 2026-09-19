"""Part 3 contracts, native artifacts, version boundaries, and rollout controls."""
import asyncio
import hashlib
import io
import json
from dataclasses import replace

import httpx
import pytest

pytest.importorskip("sqlalchemy")
pytest.importorskip("temporalio")

from docx import Document
from pypdf import PdfReader

from test_coworker import container, signed_client, account, draft
from coworker_samples import WorkModel, work_draft
from services.coworker.drafting import DeepSeekDraftModel, DraftFailure
from services.coworker.errors import CoworkerError
from services.coworker.models import Task
from services.coworker.runner import DocumentRunner
from services.coworker.schemas import TaskCreate, parse_draft
from services.coworker.skills import SKILLS
from services.coworker.work_exports import content_blocks
from services.coworker.worker import Dispatcher
from services.coworker.workflow import WorkServicesWorkflow


def enable_services(container):
    container.settings = replace(container.settings, work_services_enabled=True)
    container.repository.settings = container.settings


def create_work(container, skill_id, language="en", key="work-request-key"):
    return container.repository.create_task(account(container), TaskCreate(skill_id=skill_id,
        instruction=SKILLS[skill_id].instruction, notes="The team completed 12 reviews.", output_language=language), key)[0]


def test_catalog_auth_feature_gate_and_legacy_idempotency(container, signed_client):
    client, headers = signed_client
    assert client.get("/api/v1/skills").status_code == 401
    catalog = client.get("/api/v1/skills", headers=headers()).json()["skills"]
    assert [entry["id"] for entry in catalog] == ["report_email"]
    payload = {"instruction": "Write my email", "notes": "The team completed 12 reviews.", "document_ids": [], "output_language": "en"}
    auth = headers() | {"Idempotency-Key": "legacy-request-key"}
    first = client.post("/api/v1/tasks", headers=auth, json=payload).json()
    with container.repository.sessions() as db:
        saved = db.get(Task, first["id"])
        assert saved.fingerprint == hashlib.sha256(json.dumps(payload, sort_keys=True).encode()).hexdigest()
    replay = client.post("/api/v1/tasks", headers=auth, json=payload | {"skill_id": "report_email"})
    assert replay.status_code == 202 and replay.json()["id"] == first["id"]
    assert replay.headers["Idempotent-Replayed"] == "true"
    assert client.post("/api/v1/tasks", headers=auth, json=payload | {"skill_id": "email"}).status_code == 409
    assert client.post("/api/v1/tasks", headers=headers() | {"Idempotency-Key": "disabled-email-key"}, json=payload | {"skill_id": "email"}).status_code == 409
    assert client.post("/api/v1/tasks", headers=auth, json=payload | {"skill_id": "send_email"}).status_code == 422
    enable_services(container)
    enabled = client.get("/api/v1/skills", headers=headers()).json()["skills"]
    assert len(enabled) == 8 and all("guidance" not in entry for entry in enabled)


@pytest.mark.parametrize("skill_id,language", [("email", "ar"), ("document", "bn"), ("career", "en"), ("social", "bn"),
                                                ("meeting", "ar"), ("daily_plan", "bn"), ("personal_plan", "en")])
def test_each_service_produces_owned_native_artifacts_and_resumes(container, skill_id, language):
    enable_services(container)
    task = create_work(container, skill_id, language)
    model = WorkModel()
    runner = DocumentRunner(container, model)
    async def run():
        await runner.phase(task["id"], "extract")
        await runner.phase(task["id"], "draft")
        resumed = DocumentRunner(container, model)
        await resumed.run_for_test(task["id"])
        await resumed.run_for_test(task["id"])
    asyncio.run(run())
    result = container.repository.get_task(account(container), task["id"])
    assert result["state"] == "completed" and result["skill_id"] == skill_id
    assert result["workflow_version"] == SKILLS[skill_id].version
    assert model.calls == 1 and result["usage"]["model_attempts"] == 1
    stem = SKILLS[skill_id].filename
    assert {item["filename"] for item in result["artifacts"]} == {stem + ".docx", stem + ".pdf", stem + ".txt", "source-manifest.json"}
    for artifact in result["artifacts"]:
        owned = container.repository.artifact(account(container), artifact["id"])
        body = container.storage.get(owned["object_key"], owned["byte_size"])
        if artifact["filename"].endswith(".docx"):
            word = Document(io.BytesIO(body))
            text = "\n".join(p.text for p in word.paragraphs)
            assert "[notes]" not in text  # Provenance stays beside the ready-to-use file.
            assert any(p.text for p in word.paragraphs)
            if skill_id == "career":
                assert any(p.style.name == "List Bullet" for p in word.paragraphs)
            if language == "ar":
                assert "w:bidi" in word._element.xml and "ar" in word._element.xml
        elif artifact["filename"].endswith(".pdf"):
            assert len(PdfReader(io.BytesIO(body)).pages) > 0
        elif artifact["filename"] == "source-manifest.json":
            manifest = json.loads(body)
            assert manifest["skill_id"] == skill_id and "text" not in manifest["sources"][0]
        with pytest.raises(CoworkerError) as error:
            container.repository.artifact(account(container, "bob"), artifact["id"])
        assert error.value.status_code == 404


def test_missing_information_and_rollout_rollback_preserve_saved_work(container):
    enable_services(container)
    task = create_work(container, "email")
    asyncio.run(DocumentRunner(container, WorkModel(missing=True)).run_for_test(task["id"]))
    container.settings = replace(container.settings, work_services_enabled=False)
    container.repository.settings = container.settings
    saved = container.repository.get_task(account(container), task["id"])
    assert saved["state"] == "needs_input" and len(saved["artifacts"]) == 4
    assert create_work(container, "email")["id"] == task["id"]  # Replay is still safe with the flag off.
    with pytest.raises(CoworkerError) as error:
        create_work(container, "email", key="new-after-disable")
    assert error.value.code == "service_unavailable"


def test_complete_brief_needs_no_duplicate_notes_or_upload(container):
    enable_services(container)
    task, _ = container.repository.create_task(account(container), TaskCreate(skill_id="social",
        instruction="Draft a LinkedIn update: our team completed 12 reviews today."), "brief-only-request")
    asyncio.run(DocumentRunner(container, WorkModel()).run_for_test(task["id"]))
    result = container.repository.get_task(account(container), task["id"])
    assert result["state"] == "completed" and result["input"]["notes"] == ""
    assert [source["id"] for source in result["sources"]] == ["brief"]
    assert result["draft"]["posts"][0]["source_ids"] == ["brief"]
    with pytest.raises(ValueError):
        TaskCreate(instruction="Make a report", skill_id="report_email")


@pytest.mark.parametrize("skill_id", [key for key in SKILLS if key != "report_email"])
def test_deepseek_contract_selects_requested_service_and_rejects_wrong_output(container, skill_id):
    calls = []
    output = work_draft(skill_id, "bn")
    response_draft = output.model_dump()
    def transport(request):
        calls.append(json.loads(request.content))
        return httpx.Response(200, json={"choices": [{"finish_reason": "stop", "message": {"content": json.dumps(response_draft)}}], "usage": {"total_tokens": 123}})
    model = DeepSeekDraftModel(replace(container.settings, deepseek_api_key="test-only"), httpx.MockTransport(transport), skill_id=skill_id)
    messages = model.messages({"instruction": SKILLS[skill_id].instruction, "output_language": "bn"}, [{"id": "notes", "text": "Source data"}])
    assert SKILLS[skill_id].guidance in messages[0]["content"]
    result = asyncio.run(model.generate(messages, "bn", {"notes"}))
    assert type(result.draft) is type(output) and result.total_tokens == 123
    assert calls[0]["response_format"] == {"type": "json_object"}
    response_draft = draft("bn").model_dump()
    with pytest.raises(DraftFailure) as error:
        asyncio.run(model.generate(messages, "bn", {"notes"}))
    assert error.value.code == "invalid_draft" and error.value.total_tokens == 123


def test_bad_sources_and_service_contracts_cannot_publish_files(container):
    enable_services(container)
    task = create_work(container, "email")
    class WrongSources(WorkModel):
        async def generate(self, *args):
            value = await super().generate(*args)
            value.draft.source_ids = ["other-account-document"]
            return value
    with pytest.raises(DraftFailure):
        asyncio.run(DocumentRunner(container, WrongSources()).run_for_test(task["id"]))
    saved = container.repository.get_task(account(container), task["id"])
    assert saved["state"] == "failed" and not saved["artifacts"]
    assert saved["usage"]["accounted_tokens"] == 220
    for mutation in [{"kind": "email"}, {"output_language": "invalid-long-language"}, {"title": "bad\x00title"}]:
        with pytest.raises(ValueError):
            parse_draft("social", work_draft("social").model_dump() | mutation)
    too_many = work_draft("social").model_dump()
    too_many["posts"] *= 4
    with pytest.raises(ValueError):
        parse_draft("social", too_many)
    title, blocks = content_blocks(work_draft("meeting", "bn"))
    assert not any("None" in block.text or "Deadline:" in block.text for block in blocks)
    assert any("প্রস্তাবিত" in block.text for block in blocks)


def test_dispatcher_routes_work_services_using_persisted_version(container):
    enable_services(container)
    task = create_work(container, "social")
    class Client:
        calls = []
        async def start_workflow(self, workflow, task_id, **kwargs):
            self.calls.append((workflow, task_id, kwargs))
    client = Client()
    asyncio.run(Dispatcher(container, client).tick())
    asyncio.run(Dispatcher(container, client).tick())
    assert len(client.calls) == 1
    assert client.calls[0][0] is WorkServicesWorkflow.run and client.calls[0][1] == task["id"]
