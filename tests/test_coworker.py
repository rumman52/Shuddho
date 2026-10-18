"""Public boundary, workflow checkpoints, parser, and real artifact regressions."""
from __future__ import annotations

import asyncio
import hashlib
import io
import json
import time
import zipfile
from dataclasses import replace
from datetime import timedelta

import httpx
import pytest

pytest.importorskip("sqlalchemy", reason="Install the coworker extra for Part 2 tests")
pytest.importorskip("temporalio", reason="Install the coworker extra for Part 2 tests")

import jwt
from cryptography.hazmat.primitives.asymmetric import rsa
from docx import Document as WordDocument
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pypdf import PdfReader, PdfWriter
from sqlalchemy import select

from services.coworker.api import mount
from services.coworker.auth import JwtVerifier, Principal
from services.coworker.config import Settings, _normalize_database_url
from services.coworker.container import Container
from services.coworker.drafting import DeepSeekDraftModel, DraftFailure, DraftResult
from services.coworker.errors import CoworkerError
from services.coworker.extraction import extract_in_subprocess, extract_text
from services.coworker.exports import render_in_subprocess
from services.coworker.migrate import upgrade
from services.coworker.models import Account, Artifact, DailyUsage, ModelAttempt, Outbox, Task, utcnow
from services.coworker.runner import DocumentRunner
from services.coworker.schemas import DraftPackage, TaskCreate, UploadRequest
from coworker_worker_helpers import dispatch_tick

ISSUER = "https://identity.example.test/auth/v1"


def test_managed_postgres_url_normalization_is_secure_and_render_compatible():
    assert _normalize_database_url(
        "postgresql://user:password@db.internal:5432/app",
        "staging",
    ) == "postgresql+psycopg://user:password@db.internal:5432/app?sslmode=require"
    assert _normalize_database_url(
        "postgres://user:password@db.internal:5432/app?application_name=shuddho",
        "production",
    ) == (
        "postgresql+psycopg://user:password@db.internal:5432/app"
        "?application_name=shuddho&sslmode=require"
    )
    assert _normalize_database_url(
        "postgresql+psycopg://user:password@db:5432/app?sslmode=verify-full",
        "staging",
    ).endswith("?sslmode=verify-full")
    assert _normalize_database_url(
        "sqlite:///tmp/dev.sqlite3",
        "development",
    ) == "sqlite:///tmp/dev.sqlite3"


@pytest.fixture
def container(tmp_path):
    settings = Settings(database_url=f"sqlite:///{tmp_path / 'workspace.sqlite3'}", auth_issuer=ISSUER,
                        environment="development", storage_backend="local", local_storage_path=tmp_path / "objects")
    upgrade(settings.database_url)
    result = Container.create(settings)
    yield result
    result.repository.sessions.kw["bind"].dispose()


def principal(subject="alice"):
    return Principal(ISSUER, subject, int(time.time()) + 3600)


def account(container, subject="alice"):
    return container.repository.ensure_account(principal(subject))["account_id"]


def enable_reviewed_pilot(container):
    from services.coworker.reviewed_skill_registry import build_pilot_publication

    settings = replace(container.settings, source_revision="a" * 40)
    container.settings = settings
    container.repository.settings = settings
    container.agent.settings = settings
    worker_id = container.repository.register_worker(
        task_queue=settings.task_queue,
        source_revision=settings.source_revision,
    )
    container.repository.heartbeat_worker(worker_id)
    publication = build_pilot_publication(source_revision=settings.source_revision)
    container.repository.publish_reviewed_skill(
        publication,
        actor_reference="test-operator",
        reason="CW-11 test fixture publication",
    )
    container.repository.select_reviewed_skill(
        "research",
        publication["publication_version"],
        actor_reference="test-operator",
        reason="CW-11 test fixture selection",
        worker_source_revision=settings.source_revision,
    )


def new_task(container, owner=None, **kwargs):
    return container.repository.create_task(owner or account(container), TaskCreate(
        instruction="Prepare a project report and email.", notes="Team completed 12 reviews on 10 September 2026.", **kwargs), "task-key-" + str(time.time_ns()))[0]


def draft(language="en", missing=False):
    title, summary = {
        "en": ("Project update", "The team completed 12 reviews."),
        "bn": ("প্রকল্পের অগ্রগতি", "দলটি ১২টি পর্যালোচনা সম্পন্ন করেছে।"),
        "ar": ("تحديث المشروع", "أكمل الفريق 12 مراجعة."),
    }.get(language, ("Project update", "The team completed 12 reviews."))
    return DraftPackage.model_validate({"report": {"title": title, "summary": summary,
        "sections": [{"heading": title, "paragraphs": [summary], "source_ids": ["notes"]}]},
        "email": {"subject": title, "body": summary}, "output_language": language,
        "missing_information": ["Who is the recipient?"] if missing else []})


class FakeModel:
    def __init__(self, language="en", missing=False):
        self.calls = 0
        self.value = draft(language, missing)

    def messages(self, task, sources):
        return [{"role": "user", "content": json.dumps(sources)}]

    async def generate(self, *_args):
        self.calls += 1
        return DraftResult(self.value, 180, 5)


@pytest.fixture
def signed_client(container):
    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    jwk = jwt.algorithms.RSAAlgorithm.to_jwk(key.public_key(), as_dict=True) | {"kid": "test-key", "alg": "RS256", "use": "sig"}
    container.verifier = JwtVerifier(ISSUER, "authenticated", httpx.MockTransport(lambda request: httpx.Response(200, json={"keys": [jwk]})))
    app = FastAPI()
    mount(app, container)

    def headers(subject="alice", **changes):
        claims = {"iss": ISSUER, "sub": subject, "aud": "authenticated", "role": "authenticated",
                  "iat": int(time.time()) - 1, "exp": int(time.time()) + 3600} | changes
        token = jwt.encode(claims, key, algorithm="RS256", headers={"kid": "test-key"})
        return {"Authorization": "Bearer " + token}

    with TestClient(app) as client:
        yield client, headers


def test_simple_question_does_not_create_agent_run(container, signed_client):
    """The authenticated route rejects direct chat before durable admission."""
    from services.coworker.models import AgentOutbox, AgentRun

    settings = replace(
        container.settings, work_services_enabled=True, artifact_services_enabled=True,
        agent_runtime_enabled=True,
        intelligent_planner_enabled=True,
    )
    settings.validate()
    container.settings = settings
    container.agent.settings = settings
    account(container)
    client, headers = signed_client
    for _ in range(2):
        response = client.post(
            "/api/v1/agent-runs",
            headers={**headers(), "Idempotency-Key": "question-must-not-run"},
            json={"goal": "What is retrieval augmented generation?"},
        )
        assert response.status_code == 409, response.text
        assert response.json()["error"]["code"] == "agent_route_direct_answer"
    with container.repository.sessions() as db:
        assert db.scalars(select(AgentRun)).all() == []
        assert db.scalars(select(AgentOutbox)).all() == []
        assert db.scalars(select(Task)).all() == []




def test_server_enforced_cohort_admission_precedes_account_creation(container, signed_client):
    client, headers = signed_client
    allowed = principal("alice").account_id
    denied = principal("bob").account_id
    container.settings = replace(
        container.settings,
        cohort_enforced=True,
        cohort_account_ids=frozenset({allowed}),
        cohort_max_users=1,
    )

    admitted = client.get("/api/v1/me", headers=headers("alice"))
    assert admitted.status_code == 200
    rejected = client.get("/api/v1/me", headers=headers("bob"))
    assert rejected.status_code == 403
    assert rejected.json()["error"]["code"] == "cohort_not_enabled"

    with container.repository.sessions() as db:
        assert db.get(Account, allowed) is not None
        assert db.get(Account, denied) is None


def test_cohort_configuration_is_fail_closed(container):
    with pytest.raises(ValueError, match="requires SHUDDHO_COWORKER_COHORT_ACCOUNT_IDS"):
        replace(
            container.settings,
            cohort_enforced=True,
            cohort_account_ids=frozenset(),
        ).validate()

    valid_a = principal("cohort-a").account_id
    valid_b = principal("cohort-b").account_id
    with pytest.raises(ValueError, match="exceeds SHUDDHO_COWORKER_COHORT_MAX_USERS"):
        replace(
            container.settings,
            cohort_enforced=True,
            cohort_account_ids=frozenset({valid_a, valid_b}),
            cohort_max_users=1,
        ).validate()

    with pytest.raises(ValueError, match="64-character lowercase SHA-256"):
        replace(
            container.settings,
            cohort_enforced=True,
            cohort_account_ids=frozenset({"not-an-account-id"}),
        ).validate()

def test_auth_validation_and_private_errors(signed_client):
    client, headers = signed_client
    assert client.get("/api/v1/me").status_code == 401
    for overrides in [{"exp": 1}, {"aud": "wrong"}, {"iss": "https://other.test"}, {"role": "service_role"}, {"is_anonymous": True}]:
        assert client.get("/api/v1/me", headers=headers(**overrides)).status_code == 401
    response = client.get("/api/v1/me", headers=headers())
    assert response.status_code == 200 and response.headers["Cache-Control"] == "no-store"
    secret_notes = "PRIVATE-DOCUMENT-SENTINEL"
    invalid = client.post("/api/v1/tasks", headers=headers() | {"Idempotency-Key": "invalid-body"}, json={"notes": secret_notes})
    assert invalid.status_code == 422
    assert secret_notes not in invalid.text
    assert client.post("/api/v1/tasks", headers=headers(), content=b" " * 262145).status_code == 413
    assert client.post("/api/check", headers=headers(), json={}).status_code == 400


def test_malformed_jwks_is_a_safe_unavailable_error():
    verifier = JwtVerifier(ISSUER, "authenticated", httpx.MockTransport(lambda _: httpx.Response(200, json={"keys": [None]})))
    with pytest.raises(CoworkerError, match="temporarily unavailable") as error:
        asyncio.run(verifier._refresh())
    assert error.value.status_code == 503


def test_owned_documents_tasks_events_and_preferences(signed_client, container):
    client, headers = signed_client
    a, b = headers(), headers("bob")
    body = "Meeting notes: 12 reviews completed.".encode()
    upload = client.post("/api/v1/documents", headers=a, json={"filename": "notes.txt", "byte_size": len(body), "sha256": hashlib.sha256(body).hexdigest()})
    assert upload.status_code == 201
    doc = upload.json()
    path = f'/api/v1/documents/{doc["id"]}/content'
    assert client.put(path, headers=b, content=body).status_code == 404
    assert client.put(path, headers=a, content=b"x" * len(body)).status_code == 409
    assert client.put(path, headers=a, content=body).status_code == 200
    payload = {"instruction": "Create report and email", "document_ids": [doc["id"]], "output_language": "bn"}
    assert client.post("/api/v1/tasks", headers=b | {"Idempotency-Key": "other-user-key"}, json=payload).status_code == 404
    response = client.post("/api/v1/tasks", headers=a | {"Idempotency-Key": "submit-once"}, json=payload)
    assert response.status_code == 202
    task = response.json()
    task_path = f'/api/v1/tasks/{task["id"]}'
    replay = client.post("/api/v1/tasks", headers=a | {"Idempotency-Key": "submit-once"}, json=payload)
    assert replay.json()["id"] == task["id"] and replay.headers["Idempotent-Replayed"] == "true"
    conflict = client.post("/api/v1/tasks", headers=a | {"Idempotency-Key": "submit-once"}, json=payload | {"output_language": "en"})
    assert conflict.status_code == 409
    for suffix in ["", "/events", "/events?stream=true"]:
        assert client.get(task_path + suffix, headers=b).status_code == 404
    assert client.post(task_path + "/cancel", headers=b).status_code == 404
    assert client.get("/api/v1/tasks", headers=b).json()["tasks"] == []
    assert client.get("/api/v1/documents", headers=b).json()["documents"] == []
    client.put("/api/v1/preferences", headers=a, json={"language": "bn"})
    assert client.get("/api/v1/me", headers=b).json()["preferences"] == {}
    assert client.post(task_path + "/cancel", headers=a).json()["state"] == "cancelled"
    events = client.get(task_path + "/events?after=1", headers=a).json()
    assert events["terminal"] and events["events"][0]["sequence"] == 2
    sse = client.get(task_path + "/events?stream=true", headers=a | {"Last-Event-ID": "1"})
    assert "id: 2" in sse.text and "id: 1" not in sse.text
    assert client.get(task_path + "/events", headers=a | {"Last-Event-ID": "bad"}).status_code == 400


@pytest.mark.parametrize("language", ["en", "bn", "ar"])
def test_real_multilingual_exports_and_checkpoint_recovery(container, language):
    owner = account(container)
    task = new_task(container, owner, output_language=language)
    model = FakeModel(language)
    first_worker = DocumentRunner(container, model)
    asyncio.run(first_worker.phase(task["id"], "extract"))
    asyncio.run(first_worker.phase(task["id"], "draft"))
    # New worker instance starts from durable checkpoints. No repeat model call.
    resumed = DocumentRunner(container, model)
    asyncio.run(resumed.run_for_test(task["id"]))
    asyncio.run(DocumentRunner(container, model).run_for_test(task["id"]))
    result = container.repository.get_task(owner, task["id"])
    assert model.calls == 1 and result["state"] == "completed"
    assert all("text" not in source for source in container.repository.step(task["id"], "extract")["sources"])
    assert result["usage"] == {"model_attempts": 1, "accounted_tokens": 180}
    assert {item["filename"] for item in result["artifacts"]} == {"report.docx", "report.pdf", "email-draft.txt", "source-manifest.json"}
    for item in result["artifacts"]:
        stored = container.repository.artifact(owner, item["id"])
        raw = container.storage.get(stored["object_key"], stored["byte_size"])
        assert hashlib.sha256(raw).hexdigest() == item["sha256"]
        if item["filename"] == "report.docx":
            assert draft(language).report.title in "\n".join(p.text for p in WordDocument(io.BytesIO(raw)).paragraphs)
            if language == "ar":
                with zipfile.ZipFile(io.BytesIO(raw)) as package:
                    assert b"w:bidi" in package.read("word/document.xml")
        if item["filename"] == "report.pdf":
            assert len(PdfReader(io.BytesIO(raw)).pages) >= 1
        with pytest.raises(CoworkerError) as denied:
            container.repository.artifact(account(container, "bob"), item["id"])
        assert denied.value.status_code == 404


def test_needs_input_is_reviewable_not_a_sent_email(container, signed_client):
    client, headers = signed_client
    owner = account(container)
    task = new_task(container, owner)
    asyncio.run(DocumentRunner(container, FakeModel(missing=True)).run_for_test(task["id"]))
    result = client.get(f'/api/v1/tasks/{task["id"]}', headers=headers()).json()
    assert result["state"] == "needs_input" and result["draft"]["missing_information"]
    item = result["artifacts"][0]
    route = f'/api/v1/artifacts/{item["id"]}/download'
    assert client.get(route, headers=headers("bob")).status_code == 404
    download = client.get(route, headers=headers()).json()
    assert download["url"] is None
    assert client.get(download["content_path"], headers=headers()).status_code == 200
    assert client.get(download["content_path"], headers=headers("bob")).status_code == 404


def test_budget_unknown_outcomes_and_cancel_block_future_calls(container):
    owner = account(container)
    task = new_task(container, owner)
    repo = container.repository
    a = repo.reserve_model(task["id"], 45000)
    repo.settle_model(task["id"], a, None, None, "unknown")
    repo.settle_model(task["id"], a, 0, None, "completed")
    assert repo.get_task(owner, task["id"])["usage"]["accounted_tokens"] == 45000
    b = repo.reserve_model(task["id"], 45000)
    repo.settle_model(task["id"], b, 25, 10, "completed")
    with pytest.raises(CoworkerError) as budget:
        repo.reserve_model(task["id"], 100)
    assert budget.value.code == "task_budget"
    repo.cancel(owner, task["id"])
    with pytest.raises(CoworkerError) as cancelled:
        asyncio.run(DocumentRunner(container, FakeModel()).phase(task["id"], "draft"))
    assert cancelled.value.code == "task_cancelled"


def test_deadline_expires_without_dispatcher_and_does_not_block_quota(container):
    owner = account(container)
    task = new_task(container, owner)
    with container.repository.sessions.begin() as db:
        db.get(Task, task["id"]).deadline_at = utcnow() - timedelta(seconds=1)
    assert container.repository.get_task(owner, task["id"])["error_code"] == "task_expired"
    assert new_task(container, owner)["state"] == "queued"


def test_abandoned_upload_releases_reservation_once(container):
    owner = account(container)
    request = UploadRequest(filename="test.txt", byte_size=100, sha256="a" * 64)
    result = container.repository.create_upload(owner, request)
    from services.coworker.models import DocumentVersion
    with container.repository.sessions.begin() as db:
        db.get(DocumentVersion, result["version_id"]).expires_at = utcnow() - timedelta(minutes=3)
    assert len(container.repository.expired_uploads()) == 1
    container.repository.expired_uploads()
    with container.repository.sessions() as db:
        assert db.get(Account, owner).storage_bytes == 0
    with pytest.raises(CoworkerError):
        container.repository.finish_upload(owner, result["id"])


def test_source_deletion_checks_ownership_and_active_usage(signed_client, container):
    client, headers = signed_client
    raw = b"Project notes"
    uploaded = client.post("/api/v1/documents", headers=headers(), json={"filename": "remove.txt", "byte_size": len(raw), "sha256": hashlib.sha256(raw).hexdigest()}).json()
    route = f'/api/v1/documents/{uploaded["id"]}'
    assert client.put(route + "/content", headers=headers(), content=raw).status_code == 200
    assert client.delete(route, headers=headers("bob")).status_code == 404
    task = new_task(container, document_ids=[uploaded["id"]])
    assert client.delete(route, headers=headers()).status_code == 409
    container.repository.cancel(account(container), task["id"])
    assert client.delete(route, headers=headers()).status_code == 202
    assert client.get("/api/v1/documents", headers=headers()).json()["documents"] == []
    assert client.put(route + "/content", headers=headers(), content=raw).status_code == 404
    # A cancelled task can still be consumed from the outbox after source deletion.
    assert container.repository.worker_task(task["id"], False)["state"] == "cancelled"


def test_source_scans_expansion_and_production_tls():
    from pypdf.generic import DictionaryObject, NameObject, NumberObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(100, 100)
    image = DecodedStreamObject()
    image.set_data(b"\xff\xff\xff")
    image.update({NameObject("/Type"): NameObject("/XObject"), NameObject("/Subtype"): NameObject("/Image"),
                  NameObject("/Width"): NumberObject(1), NameObject("/Height"): NumberObject(1),
                  NameObject("/ColorSpace"): NameObject("/DeviceRGB"), NameObject("/BitsPerComponent"): NumberObject(8)})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/XObject"): DictionaryObject({NameObject("/Image1"): writer._add_object(image)})})
    raw = io.BytesIO()
    writer.write(raw)
    with pytest.raises(CoworkerError) as scan:
        extract_in_subprocess(raw.getvalue(), "pdf", 20000)
    assert scan.value.code == "ocr_required"
    expanded = io.BytesIO()
    with zipfile.ZipFile(expanded, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b" " * (33 * 1024 * 1024))
    with pytest.raises(CoworkerError) as bomb:
        extract_text(expanded.getvalue(), "docx")
    assert bomb.value.code == "document_expansion_limit"
    settings = Settings(database_url="postgresql+psycopg://user:placeholder@host/db", auth_issuer=ISSUER, storage_bucket="test")
    with pytest.raises(ValueError, match="sslmode"):
        settings.validate()


def test_unexpected_database_errors_do_not_echo_private_parameters(signed_client, container, monkeypatch, caplog):
    client, headers = signed_client
    assert client.get("/api/v1/me", headers=headers()).status_code == 200
    def fail(_owner):
        raise RuntimeError("PRIVATE-DATABASE-PARAMETER")
    monkeypatch.setattr(container.repository, "list_tasks", fail)
    response = client.get("/api/v1/tasks", headers=headers())
    assert response.status_code == 503
    assert "PRIVATE-DATABASE-PARAMETER" not in response.text + caplog.text


def test_file_parser_limits_and_real_docx_pdf():
    assert extract_in_subprocess("বাংলা notes".encode(), "txt", 20000) == "বাংলা notes"
    for data, code in [(b"", "empty_document"), (b"\xff", "text_encoding"), (b"\x00binary", "document_invalid")]:
        with pytest.raises(CoworkerError) as failure:
            extract_in_subprocess(data, "txt", 20000)
        assert failure.value.code == code
    doc = WordDocument()
    doc.add_paragraph("Project facts")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "12 reviews"
    buf = io.BytesIO()
    doc.save(buf)
    assert "12 reviews" in extract_in_subprocess(buf.getvalue(), "docx", 20000)
    dangerous = io.BytesIO()
    with zipfile.ZipFile(dangerous, "w") as package:
        package.writestr("word/vbaProject.bin", "macro")
        package.writestr("word/document.xml", "<x/>")
    with pytest.raises(CoworkerError) as failure:
        extract_in_subprocess(dangerous.getvalue(), "docx", 20000)
    assert failure.value.code == "document_active_content"
    writer = PdfWriter()
    for _ in range(41):
        writer.add_blank_page(100, 100)
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(CoworkerError) as failure:
        extract_text(buffer.getvalue(), "pdf")
    assert failure.value.code == "page_limit"


@pytest.mark.parametrize("change", ["truncated", "source", "language", "tool", "control"])
def test_model_rejects_unusable_or_unattributed_output(container, change):
    payload = draft().model_dump()
    if change == "source":
        payload["report"]["sections"][0]["source_ids"] = ["invented"]
    if change == "language":
        payload["output_language"] = "fr"
    if change == "control":
        payload["report"]["title"] = "bad\x00title"
    message = {"content": json.dumps(payload)}
    if change == "tool":
        message["tool_calls"] = [{"name": "send_email"}]
    envelope = {"choices": [{"finish_reason": "length" if change == "truncated" else "stop", "message": message}], "usage": {"total_tokens": 123}}
    model = DeepSeekDraftModel(replace(container.settings, deepseek_api_key="test-only-placeholder"), httpx.MockTransport(lambda _: httpx.Response(200, json=envelope)))
    with pytest.raises(DraftFailure) as failure:
        asyncio.run(model.generate([], "en", {"notes"}))
    assert failure.value.code == "invalid_draft" and failure.value.total_tokens == 123


def test_model_success_and_retryable_provider_failure(container):
    settings = replace(container.settings, deepseek_api_key="test-only-placeholder")
    def respond(request):
        body = json.loads(request.content)
        assert body["thinking"] == {"type": "disabled"}
        assert body["response_format"] == {"type": "json_object"}
        return httpx.Response(200, json={"choices": [{"finish_reason": "stop", "message": {"content": draft().model_dump_json()}}], "usage": {"total_tokens": 200}})
    model = DeepSeekDraftModel(settings, httpx.MockTransport(respond))
    assert asyncio.run(model.generate([], "en", {"notes"})).total_tokens == 200
    model = DeepSeekDraftModel(settings, httpx.MockTransport(lambda _: httpx.Response(429, text="PRIVATE PROVIDER BODY")))
    with pytest.raises(DraftFailure) as failure:
        asyncio.run(model.generate([], "en", {"notes"}))
    assert failure.value.retryable and "PRIVATE" not in failure.value.message


def test_outbox_recovery_does_not_start_another_workflow(container):
    from temporalio.exceptions import WorkflowAlreadyStartedError
    task = new_task(container)
    started = set()
    class Client:
        async def start_workflow(self, _workflow, task_id, **options):
            assert options["id"] == "shuddho-task-" + task_id
            if task_id in started:
                raise WorkflowAlreadyStartedError(options["id"], "shuddho_report_email_v1")
            started.add(task_id)
    asyncio.run(dispatch_tick(container, Client()))
    with container.repository.sessions.begin() as db:
        row = db.get(Outbox, task["id"])
        row.delivered = False
        row.lease_until = utcnow() - timedelta(seconds=1)
    asyncio.run(dispatch_tick(container, Client()))
    assert started == {task["id"]}
    with container.repository.sessions() as db:
        assert db.get(Outbox, task["id"]).delivered



def test_agent_runtime_is_fail_closed_by_default(signed_client):
    client, headers = signed_client
    tools = client.get("/api/v1/agent-tools", headers=headers())
    assert tools.status_code == 200
    assert tools.json() == {"enabled": False, "tools": []}
    response = client.post("/api/v1/agent-runs", headers=headers() | {"Idempotency-Key": "agent-disabled"},
                           json={"goal": "Prepare an investor meeting pack.", "output_language": "en"})
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "agent_runtime_unavailable"


def test_agent_runtime_owned_idempotent_and_cancelable(signed_client, container):
    from services.coworker.agent_tools import TOOLS
    client, headers = signed_client
    enabled = replace(container.settings, agent_runtime_enabled=True, work_services_enabled=True,
                      artifact_services_enabled=True, research_services_enabled=False, actions_enabled=False)
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled

    catalog = client.get("/api/v1/agent-tools", headers=headers()).json()
    names = {item["name"] for item in catalog["tools"]}
    assert catalog["enabled"] is True
    assert {"document.create", "email.draft", "meeting.prepare", "presentation.create", "spreadsheet.create"} <= names
    assert "research.search" not in names and "email.send" not in names and "calendar.create" not in names
    assert TOOLS["email.send"].consequential and TOOLS["email.send"].approval_required
    assert TOOLS["calendar.create"].consequential and TOOLS["calendar.create"].approval_required

    payload = {"goal": "Prepare an investor meeting pack from my current sources.", "output_language": "en"}
    first = client.post("/api/v1/agent-runs", headers=headers() | {"Idempotency-Key": "agent-run-once"}, json=payload)
    assert first.status_code == 202
    run = first.json()
    assert first.headers["Idempotent-Replayed"] == "false"
    assert run["state"] == "queued" and run["phase"] == "planning"
    assert run["steps"][0]["ordinal"] == 0 and run["steps"][0]["tool"] is None
    assert run["tool_invocations"] == []

    replay = client.post("/api/v1/agent-runs", headers=headers() | {"Idempotency-Key": "agent-run-once"}, json=payload)
    assert replay.status_code == 202
    assert replay.json()["id"] == run["id"]
    assert replay.headers["Idempotent-Replayed"] == "true"
    conflict = client.post("/api/v1/agent-runs", headers=headers() | {"Idempotency-Key": "agent-run-once"},
                           json=payload | {"goal": "Prepare a different goal."})
    assert conflict.status_code == 409
    assert conflict.json()["error"]["code"] == "idempotency_conflict"

    assert client.get(f'/api/v1/agent-runs/{run["id"]}', headers=headers("bob")).status_code == 404
    assert client.post(f'/api/v1/agent-runs/{run["id"]}/cancel', headers=headers("bob")).status_code == 404
    cancelled = client.post(f'/api/v1/agent-runs/{run["id"]}/cancel', headers=headers())
    assert cancelled.status_code == 200
    assert cancelled.json()["state"] == "cancelled"
    assert cancelled.json()["steps"][0]["state"] == "cancelled"
    assert client.get("/api/v1/agent-runs", headers=headers()).json()["runs"][0]["id"] == run["id"]


def test_agent_runtime_requires_owned_uploaded_documents(signed_client, container):
    client, headers = signed_client
    enabled = replace(container.settings, agent_runtime_enabled=True)
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    missing = "11111111-1111-1111-1111-111111111111"
    response = client.post("/api/v1/agent-runs", headers=headers() | {"Idempotency-Key": "agent-missing-doc"},
                           json={"goal": "Summarize this source.", "document_ids": [missing], "output_language": "en"})
    assert response.status_code == 404



def test_agent_plan_is_bounded_registered_and_source_scoped(signed_client, container):
    from services.coworker.agent_schemas import AgentPlanStep, AgentRunCreate
    client, headers = signed_client
    enabled = replace(container.settings, agent_runtime_enabled=True, work_services_enabled=True, max_active_agent_runs=10)
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    owner = account(container)

    run, _ = container.agent.create(owner, AgentRunCreate(
        goal="Prepare a professional follow-up email.", output_language="en"
    ), "agent-plan-ledger")
    events = client.get(f'/api/v1/agent-runs/{run["id"]}/events', headers=headers()).json()
    assert events["sequence"] == 1 and events["events"][0]["phase"] == "planning"

    planned = container.agent.save_plan(owner, run["id"], [
        AgentPlanStep(tool="email.draft", arguments={
            "instruction": "Draft a concise follow-up email.",
            "notes": "Thank the team for the project review.",
            "document_ids": [],
            "output_language": "en",
        }),
    ])
    assert planned["phase"] == "planned"
    assert planned["event_sequence"] == 2
    assert planned["steps"][0]["tool"] == "email.draft"
    invocation = planned["tool_invocations"][0]
    assert invocation["tool"] == "email.draft"
    assert invocation["state"] == "prepared"
    assert invocation["consequential"] is False
    assert invocation["approval_required"] is False
    assert invocation["receipt"] is None

    after = client.get(f'/api/v1/agent-runs/{run["id"]}/events?after=1', headers=headers()).json()
    assert [item["sequence"] for item in after["events"]] == [2]
    assert after["events"][0]["phase"] == "planned"
    assert client.get(f'/api/v1/agent-runs/{run["id"]}/events', headers=headers("bob")).status_code == 404

    with pytest.raises(CoworkerError) as repeated:
        container.agent.save_plan(owner, run["id"], [AgentPlanStep(tool="email.draft", arguments={
            "instruction": "Draft another email.", "notes": "", "document_ids": [], "output_language": "en",
        })])
    assert repeated.value.code == "plan_already_saved"

    other, _ = container.agent.create(owner, AgentRunCreate(
        goal="Use only the sources attached to this run.", output_language="en"
    ), "agent-plan-source-scope")
    with pytest.raises(CoworkerError) as scope:
        container.agent.save_plan(owner, other["id"], [AgentPlanStep(tool="document.create", arguments={
            "instruction": "Create a memo.",
            "notes": "",
            "document_ids": ["11111111-1111-1111-1111-111111111111"],
            "output_language": "en",
        })])
    assert scope.value.code == "tool_source_scope"

    third, _ = container.agent.create(owner, AgentRunCreate(
        goal="Reject model-invented tools.", output_language="en"
    ), "agent-plan-tool-scope")
    with pytest.raises(CoworkerError) as unknown:
        container.agent.save_plan(owner, third["id"], [AgentPlanStep(tool="shell.run", arguments={})])
    assert unknown.value.code == "tool_not_supported"



def test_agent_task_scope_survives_plan_replan_and_worker_claim(container):
    """A durable derived-run grant must be enforced past the planner boundary."""
    from services.coworker.agent_runtime import AgentRuntime
    from services.coworker.agent_schemas import AgentPlanStep, AgentRunCreate
    from services.coworker.models import AgentStep, BrowserSession, SandboxSession, ToolInvocation

    settings = replace(
        container.settings, agent_runtime_enabled=True, work_services_enabled=True,
        intelligent_planner_enabled=True,
    )
    container.settings = settings
    container.repository.settings = settings
    container.agent.settings = settings
    owner = account(container)
    run, _ = container.agent.create(
        owner, AgentRunCreate(goal="Prepare a daily plan.", output_language="en"),
        "run-scope-worker-claim", tool_allowlist=["daily_plan.create"],
    )
    arguments = {
        "instruction": "Prepare a daily plan.", "notes": "",
        "document_ids": [], "output_language": "en",
    }
    forbidden = AgentPlanStep(tool="email.draft", arguments=arguments)
    forbidden_browser = AgentPlanStep(tool="browser.read", arguments={})
    allowed = AgentPlanStep(tool="daily_plan.create", arguments=arguments)

    # A deterministic planner or direct replan must not bypass the stored scope.
    for disallowed in (forbidden, forbidden_browser):
        with pytest.raises(CoworkerError) as initial:
            container.agent.save_plan(owner, run["id"], [disallowed])
        assert initial.value.code == "agent_tool_scope"
    with container.repository.sessions() as db:
        assert db.scalar(select(ToolInvocation.id).where(
            ToolInvocation.run_id == run["id"],
        )) is None

    saved = container.agent.save_plan(owner, run["id"], [allowed])
    assert saved["tool_invocations"][0]["tool"] == "daily_plan.create"
    for disallowed in (forbidden, forbidden_browser):
        with pytest.raises(CoworkerError) as replanned:
            container.agent.replace_remaining_plan(owner, run["id"], 1, [disallowed])
        assert replanned.value.code == "agent_tool_scope"

    # Simulate a stale or altered prepared invocation after planning. The
    # actual worker must reject it before any task, provider or artifact work.
    with container.repository.sessions.begin() as db:
        invocation = db.scalar(select(ToolInvocation).where(
            ToolInvocation.run_id == run["id"],
        ))
        step = db.get(AgentStep, invocation.step_id)
        assert invocation.tool_name == step.tool_name == "daily_plan.create"
        invocation.tool_name = "browser.read"
        step.tool_name = "browser.read"

    runtime = AgentRuntime(container, DocumentRunner(container, FakeModel()))
    with pytest.raises(CoworkerError) as worker:
        asyncio.run(runtime.execute_step(run["id"], 1))
    assert worker.value.code == "agent_tool_scope"
    with container.repository.sessions() as db:
        assert db.scalar(select(Task.id).where(Task.agent_run_id == run["id"])) is None
        invocation = db.scalar(select(ToolInvocation).where(
            ToolInvocation.run_id == run["id"],
        ))
        assert invocation.state == "prepared"
        assert db.scalar(select(BrowserSession.id).where(
            BrowserSession.owner_id == owner,
        )) is None
        assert db.scalar(select(SandboxSession.id).where(
            SandboxSession.owner_id == owner,
        )) is None



def test_reviewed_agent_research_requires_exact_worker_grant_and_version(container, monkeypatch):
    """CW-11 opt-in worker gate binds reviewed research to SQL-owned run and tool."""
    from services.coworker.agent_schemas import AgentPlanStep, AgentRunCreate
    from services.coworker.models import AgentRun, Step, ToolInvocation
    from services.coworker.schemas import ResearchOptions
    from services.coworker.skill_manifests import pilot_skill_manifest
    from services.coworker.skills import SKILLS

    with pytest.raises(ValueError, match="Reviewed Agent research requires"):
        replace(
            container.settings, agent_reviewed_research_skill_enabled=True,
        ).validate()
    settings = replace(
        container.settings, agent_runtime_enabled=True,
        research_services_enabled=True, agent_reviewed_research_skill_enabled=True,
    )
    container.settings = settings
    container.repository.settings = settings
    container.agent.settings = settings
    enable_reviewed_pilot(container)
    owner = account(container)
    outsider = account(container, "bob")
    run, _ = container.agent.create(
        owner, AgentRunCreate(goal="Research the current topic.", output_language="en"),
        "reviewed-research-run", tool_allowlist=["research.search"],
    )
    arguments = {
        "instruction": "Research the current topic.", "notes": "",
        "document_ids": [], "output_language": "en",
        "query": "reliable public sources", "time_range": "any",
    }
    container.agent.save_plan(owner, run["id"], [
        AgentPlanStep(tool="research.search", arguments=arguments),
    ])
    with container.repository.sessions() as db:
        invocation = db.scalar(select(ToolInvocation).where(
            ToolInvocation.run_id == run["id"],
        ))
        step_id = invocation.step_id

    request = TaskCreate(
        skill_id="research", instruction=arguments["instruction"], notes="",
        research=ResearchOptions(query=arguments["query"], time_range="any"),
    )
    with pytest.raises(CoworkerError) as cross_owner:
        container.repository.create_task(
            outsider, request, "wrong-owner-research", enqueue=False,
            agent_run_id=run["id"], agent_step_id=step_id,
        )
    assert cross_owner.value.code == "skill_run_scope"

    with pytest.raises(CoworkerError) as changed_query:
        container.repository.create_task(
            owner,
            TaskCreate(
                skill_id="research", instruction=arguments["instruction"],
                research=ResearchOptions(query="a different query", time_range="any"),
            ),
            "changed-query-research", enqueue=False,
            agent_run_id=run["id"], agent_step_id=step_id,
        )
    assert changed_query.value.code == "skill_run_scope"

    # Public callers can predict the agent:<run>:<ordinal> key. Reusing a
    # pre-created standalone Task must not smuggle it into this Agent step.
    collision_key = f"agent:{run['id']}:1"
    standalone_collision, _ = container.repository.create_task(
        owner, request, collision_key, enqueue=False,
    )
    with pytest.raises(CoworkerError) as collision:
        container.repository.create_task(
            owner, request, collision_key, enqueue=False,
            agent_run_id=run["id"], agent_step_id=step_id,
        )
    assert collision.value.code == "task_agent_binding_conflict"
    assert standalone_collision["id"] != run["id"]

    task, created = container.repository.create_task(
        owner, request, "owned-reviewed-research", enqueue=False,
        agent_run_id=run["id"], agent_step_id=step_id,
    )
    assert created and task["workflow_version"] == "work_research_reviewed_v2"
    with pytest.raises(CoworkerError) as duplicate_invocation:
        container.repository.create_task(
            owner, request, "second-task-same-agent-invocation", enqueue=False,
            agent_run_id=run["id"], agent_step_id=step_id,
        )
    assert duplicate_invocation.value.code == "task_agent_binding_conflict"
    # An older worker does not recognize this distinct workflow identity,
    # so rolling it back cannot silently run reviewed tasks as legacy research.
    from services.coworker.skills import skill_for_version
    with pytest.raises(CoworkerError) as old_reader:
        skill_for_version(task["workflow_version"])
    assert old_reader.value.code == "workflow_version"
    assert container.repository.worker_task(task["id"], False)["agent_step_id"] == step_id
    with pytest.raises(CoworkerError) as reverse_collision:
        container.repository.create_task(
            owner, request, "owned-reviewed-research", enqueue=False,
        )
    assert reverse_collision.value.code == "task_agent_binding_conflict"
    with container.repository.sessions() as db:
        pinned = db.get(Step, (task["id"], "reviewed_skill_identity")).output
    assert pinned["review_sha256"] == pilot_skill_manifest().review_sha256()
    assert pinned["publication_version"] == "research_reviewed_v1"
    assert pinned["task_workflow_version"] == "work_research_reviewed_v2"
    assert pinned["source_revision"] == "a" * 40

    # A stored review SHA is tied to the admitted Task: changed bytes
    # cannot be interpreted as an authorized publication/revision.
    with container.repository.sessions.begin() as db:
        row = db.get(Step, (task["id"], "reviewed_skill_identity"))
        row.output = row.output | {"review_sha256": "0" * 64}
    with pytest.raises(CoworkerError) as changed_review:
        container.repository.worker_task(task["id"], False)
    assert changed_review.value.code == "skill_manifest_unreviewed"
    with container.repository.sessions.begin() as db:
        db.get(Step, (task["id"], "reviewed_skill_identity")).output = pinned

    with container.repository.sessions.begin() as db:
        db.get(Task, task["id"]).workflow_version = "work_email_v1"
    with pytest.raises(CoworkerError) as changed_skill:
        container.repository.worker_task(task["id"], False)
    assert changed_skill.value.code == "skill_run_scope"
    with container.repository.sessions.begin() as db:
        db.get(Task, task["id"]).workflow_version = "work_research_reviewed_v2"

    with container.repository.sessions.begin() as db:
        db.scalar(select(ToolInvocation).where(
            ToolInvocation.run_id == run["id"],
        )).tool_version = "unreviewed"
    with pytest.raises(CoworkerError) as changed_tool:
        container.repository.worker_task(task["id"], False)
    assert changed_tool.value.code == "skill_run_scope"

    with container.repository.sessions.begin() as db:
        db.scalar(select(ToolInvocation).where(
            ToolInvocation.run_id == run["id"],
        )).tool_version = "1"
        db.get(AgentRun, run["id"]).tool_allowlist = []
    with pytest.raises(CoworkerError) as revoked:
        container.repository.worker_task(task["id"], False)
    assert revoked.value.code == "skill_manifest_unreviewed"

    with container.repository.sessions.begin() as db:
        db.get(AgentRun, run["id"]).tool_allowlist = ["research.search"]
    with monkeypatch.context() as patch:
        patch.setitem(SKILLS, "research", replace(
            SKILLS["research"], guidance=SKILLS["research"].guidance + " unreviewed",
        ))
        with pytest.raises(CoworkerError) as policy_drift:
            container.repository.worker_task(task["id"], False)
        assert policy_drift.value.code == "skill_manifest_unreviewed"

    with container.repository.sessions.begin() as db:
        db.get(AgentRun, run["id"]).deadline_at = utcnow() - timedelta(seconds=1)
    with pytest.raises(CoworkerError) as expired:
        container.repository.worker_task(task["id"], False)
    assert expired.value.code == "skill_run_scope"
    with container.repository.sessions.begin() as db:
        db.get(AgentRun, run["id"]).deadline_at = utcnow() + timedelta(minutes=20)

    with container.repository.sessions.begin() as db:
        db.get(AgentRun, run["id"]).cancel_requested = True
    with pytest.raises(CoworkerError) as cancelled:
        container.repository.worker_task(task["id"], False)
    assert cancelled.value.code == "skill_run_scope"
    # The reservation is the paid-provider boundary. A cancellation after
    # task admission must still prevent the credit reservation and search.
    with pytest.raises(CoworkerError) as no_search:
        container.repository.reserve_search(task["id"])
    assert no_search.value.code == "skill_run_scope"
    assert container.repository.step(task["id"], "search_attempt") is None
    with pytest.raises(CoworkerError) as cancelled_replay:
        container.repository.create_task(
            owner, request, "owned-reviewed-research", enqueue=False,
            agent_run_id=run["id"], agent_step_id=step_id,
        )
    assert cancelled_replay.value.code == "skill_run_scope"

    # The owning account has two active tasks from the collision and reviewed
    # admission cases. Verify standalone compatibility in a separate account
    # so the regression does not trip the independent active-task quota.
    standalone, _ = container.repository.create_task(
        outsider, request, "standalone-research", enqueue=False,
    )
    assert container.repository.worker_task(standalone["id"], False)["agent_run_id"] is None




def test_reviewed_research_gate_disable_preserves_pinned_task(container):
    """Disabling the admission flag must not downgrade an already pinned reviewed Task."""
    from services.coworker.agent_schemas import AgentPlanStep, AgentRunCreate
    from services.coworker.models import Step
    from services.coworker.schemas import ResearchOptions

    enabled = replace(
        container.settings,
        agent_runtime_enabled=True,
        research_services_enabled=True,
        agent_reviewed_research_skill_enabled=True,
    )
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    enable_reviewed_pilot(container)

    owner = account(container)
    run, _ = container.agent.create(
        owner, AgentRunCreate(goal="Review public evidence.", output_language="en"),
        "reviewed-rollback-run", tool_allowlist=["research.search"],
    )
    plan = container.agent.save_plan(owner, run["id"], [AgentPlanStep(
        tool="research.search", arguments={
            "instruction": "Review public evidence.", "notes": "",
            "document_ids": [], "output_language": "en",
            "query": "reviewed public sources", "time_range": "any",
        },
    )])
    request = TaskCreate(
        skill_id="research", instruction="Review public evidence.",
        research=ResearchOptions(query="reviewed public sources"),
    )
    task, _ = container.repository.create_task(
        owner, request, "reviewed-rollback-task", enqueue=False,
        agent_run_id=run["id"],
        agent_step_id=plan["tool_invocations"][0]["step_id"],
    )
    assert task["workflow_version"] == "work_research_reviewed_v2"
    with container.repository.sessions() as db:
        assert db.get(Step, (task["id"], "reviewed_skill_identity")) is not None

    # The flag controls new rollout admission, not the immutable identity of
    # a Task that was already accepted under the selected publication.
    disabled = replace(
        container.settings,
        agent_reviewed_research_skill_enabled=False,
    )
    container.settings = disabled
    container.repository.settings = disabled
    container.agent.settings = disabled

    pinned = container.repository.worker_task(task["id"], False)
    assert pinned["agent_run_id"] == run["id"]
    assert pinned["workflow_version"] == "work_research_reviewed_v2"

    replay, created = container.repository.create_task(
        owner, request, "reviewed-rollback-task", enqueue=False,
        agent_run_id=run["id"],
        agent_step_id=plan["tool_invocations"][0]["step_id"],
    )
    assert not created
    assert replay["id"] == task["id"]
    assert replay["workflow_version"] == "work_research_reviewed_v2"

    # Old, unpinned standalone research remains legacy; disabling the reviewed
    # admission flag must not retroactively enroll it into the pilot.
    legacy_owner = account(container, "legacy-after-rollback")
    legacy, _ = container.repository.create_task(
        legacy_owner, request, "legacy-standalone-research", enqueue=False,
    )
    assert container.repository.worker_task(legacy["id"], False)["agent_run_id"] is None
    with container.repository.sessions() as db:
        assert db.get(Step, (legacy["id"], "reviewed_skill_identity")) is None

    # Terminal retries return their committed result without new work even
    # while the reviewed admission flag is disabled.
    with container.repository.sessions.begin() as db:
        db.get(Task, task["id"]).state = "completed"
    terminal, created = container.repository.create_task(
        owner, request, "reviewed-rollback-task", enqueue=False,
        agent_run_id=run["id"],
        agent_step_id=plan["tool_invocations"][0]["step_id"],
    )
    assert not created
    assert terminal["id"] == task["id"]
    assert terminal["state"] == "completed"


def test_reviewed_research_rechecks_revocation_after_search_reservation(container, monkeypatch):
    """No provider call after an Agent is cancelled during the reserved/read gap."""
    from services.coworker.agent_schemas import AgentPlanStep, AgentRunCreate
    from services.coworker.models import AgentRun
    from services.coworker.schemas import ResearchOptions

    settings = replace(
        container.settings, agent_runtime_enabled=True,
        research_services_enabled=True, agent_reviewed_research_skill_enabled=True,
    )
    container.settings = settings
    container.repository.settings = settings
    container.agent.settings = settings
    enable_reviewed_pilot(container)
    owner = account(container)
    run, _ = container.agent.create(
        owner, AgentRunCreate(goal="Research public evidence.", output_language="en"),
        "research-reserve-race", tool_allowlist=["research.search"],
    )
    plan = container.agent.save_plan(owner, run["id"], [AgentPlanStep(
        tool="research.search", arguments={
            "instruction": "Research public evidence.", "notes": "",
            "document_ids": [], "output_language": "en",
            "query": "reliable public sources", "time_range": "any",
        },
    )])
    step_id = plan["tool_invocations"][0]["step_id"]
    task, _ = container.repository.create_task(
        owner,
        TaskCreate(
            skill_id="research", instruction="Research public evidence.",
            research=ResearchOptions(query="reliable public sources"),
        ),
        "reviewed-race-task", enqueue=False,
        agent_run_id=run["id"], agent_step_id=step_id,
    )
    worker_task = container.repository.worker_task(task["id"], False)

    class NeverSearch:
        calls = 0

        async def retrieve(self, _options):
            self.calls += 1
            raise AssertionError("Search must not reach the provider after revocation")

    provider = NeverSearch()
    runner = DocumentRunner(container, FakeModel(), research=provider)
    reserve = container.repository.reserve_search

    def revoke_immediately_after_reserving(task_id):
        reserve(task_id)
        with container.repository.sessions.begin() as db:
            db.get(AgentRun, run["id"]).cancel_requested = True

    monkeypatch.setattr(
        container.repository, "reserve_search", revoke_immediately_after_reserving,
    )
    with pytest.raises(CoworkerError) as denied:
        asyncio.run(runner.research(worker_task))
    assert denied.value.code == "skill_run_scope"
    assert provider.calls == 0
    receipt = container.repository.step(task["id"], "search_attempt")
    assert receipt["state"] == "failed"
    assert receipt["accounted_credits"] == 0



def test_reviewed_research_rejects_new_document_version_under_same_id(container):
    """Agent research cannot silently upgrade the exact source version it admitted."""
    from uuid import UUID, uuid4
    from services.coworker.agent_schemas import AgentPlanStep, AgentRunCreate
    from services.coworker.models import DocumentVersion
    from services.coworker.schemas import ResearchOptions

    settings = replace(
        container.settings, agent_runtime_enabled=True,
        research_services_enabled=True, agent_reviewed_research_skill_enabled=True,
    )
    container.settings = settings
    container.repository.settings = settings
    container.agent.settings = settings
    enable_reviewed_pilot(container)
    owner = account(container)
    raw = b"Reviewed document version"
    uploaded = container.repository.create_upload(
        owner, UploadRequest(
            filename="reviewed.txt", byte_size=len(raw),
            sha256=hashlib.sha256(raw).hexdigest(),
        ),
    )
    container.repository.finish_upload(owner, uploaded["id"])
    document_id = UUID(uploaded["id"])
    run, _ = container.agent.create(
        owner,
        AgentRunCreate(
            goal="Research a public comparison report.",
            document_ids=[document_id], output_language="en",
        ),
        "reviewed-source-version-run", tool_allowlist=["research.search"],
    )
    plan = container.agent.save_plan(owner, run["id"], [AgentPlanStep(
        tool="research.search", arguments={
            "instruction": "Research a public comparison report.",
            "notes": "", "document_ids": [str(document_id)],
            "output_language": "en", "query": "public source comparison",
            "time_range": "any",
        },
    )])
    task, _ = container.repository.create_task(
        owner,
        TaskCreate(
            skill_id="research",
            instruction="Research a public comparison report.",
            document_ids=[document_id],
            research=ResearchOptions(query="public source comparison"),
        ),
        "reviewed-source-version-task", enqueue=False,
        agent_run_id=run["id"],
        agent_step_id=plan["tool_invocations"][0]["step_id"],
    )
    assert container.repository.worker_task(task["id"], False)["documents"][0]["version_id"] == uploaded["version_id"]

    with container.repository.sessions.begin() as db:
        old = db.get(DocumentVersion, uploaded["version_id"])
        later = DocumentVersion(
            id=str(uuid4()), version=2, document_id=old.document_id,
            owner_id=owner, kind=old.kind, byte_size=old.byte_size,
            sha256=old.sha256, object_key=old.object_key + "-next",
            state="uploaded", expires_at=old.expires_at,
        )
        db.add(later)
        db.get(Task, task["id"]).input_versions = [later.id]
    with pytest.raises(CoworkerError) as changed_version:
        container.repository.worker_task(task["id"], False)
    assert changed_version.value.code == "skill_run_scope"


def test_reviewed_agent_research_rejects_ungranted_planner_step_before_task_admission(container):
    """No task may be created using a different legitimate run step as research authority."""
    from services.coworker.agent_schemas import AgentPlanStep, AgentRunCreate
    from services.coworker.schemas import ResearchOptions

    settings = replace(
        container.settings, agent_runtime_enabled=True, research_services_enabled=True,
        work_services_enabled=True, agent_reviewed_research_skill_enabled=True,
    )
    container.settings = settings
    container.repository.settings = settings
    container.agent.settings = settings
    enable_reviewed_pilot(container)
    owner = account(container)
    run, _ = container.agent.create(
        owner, AgentRunCreate(goal="Prepare a daily plan.", output_language="en"),
        "ungranted-research-run", tool_allowlist=["daily_plan.create"],
    )
    container.agent.save_plan(owner, run["id"], [AgentPlanStep(
        tool="daily_plan.create", arguments={
            "instruction": "Prepare a daily plan.", "notes": "",
            "document_ids": [], "output_language": "en",
        },
    )])
    with container.repository.sessions() as db:
        from services.coworker.models import ToolInvocation
        step_id = db.scalar(select(ToolInvocation).where(
            ToolInvocation.run_id == run["id"],
        )).step_id
    with pytest.raises(CoworkerError) as denied:
        container.repository.create_task(
            owner,
            TaskCreate(
                skill_id="research", instruction="Research a topic.",
                research=ResearchOptions(query="public topic", time_range="any"),
            ),
            "research-bypass-task", enqueue=False,
            agent_run_id=run["id"], agent_step_id=step_id,
        )
    assert denied.value.code == "skill_run_scope"
    with container.repository.sessions() as db:
        assert db.scalar(select(Task.id).where(
            Task.agent_run_id == run["id"],
        )) is None


def test_restricted_run_planner_and_fallback_never_offer_extra_tools(container):
    """Initial planning, replan and provider-error fallback share the SQL-owned grant."""
    from services.coworker.agent_runtime import AgentRuntime
    from services.coworker.agent_schemas import AgentPlannerProposal, AgentRunCreate
    from services.coworker.agent_planner import deterministic_plan
    from services.coworker.agent_planning_model import PlannerFailure

    settings = replace(
        container.settings, agent_runtime_enabled=True, intelligent_planner_enabled=True,
        work_services_enabled=True, max_agent_planner_calls=3,
        agent_planner_token_budget=16000,
    )
    container.settings = settings
    container.repository.settings = settings
    container.agent.settings = settings
    owner = account(container)
    scope = ["daily_plan.create"]
    goal = "Prepare a daily plan and follow-up email."
    run, _ = container.agent.create(
        owner, AgentRunCreate(goal=goal, output_language="en"),
        "restricted-planner-grant", tool_allowlist=scope,
    )
    seen = []

    class Planner:
        async def propose(self, _goal, tools, **kwargs):
            seen.append((list(tools), kwargs["reason"]))
            return AgentPlannerProposal.model_validate({
                "steps": [{"tool": "daily_plan.create", "objective": "Create a plan."}],
            }), 10, 1

    runtime = AgentRuntime(container, DocumentRunner(container, FakeModel()), planner=Planner())
    assert asyncio.run(runtime.plan_for_worker(run["id"])) == 1
    assert asyncio.run(runtime.replan(run["id"], 1)) == 1
    assert seen == [(scope, "initial"), (scope, "capability_changed")]
    saved = container.agent.get(owner, run["id"])
    assert [step["tool"] for step in saved["tool_invocations"]] == scope

    fallback, _ = container.agent.create(
        owner, AgentRunCreate(goal=goal, output_language="en"),
        "restricted-planner-fallback", tool_allowlist=scope,
    )

    class UnavailablePlanner:
        async def propose(self, _goal, tools, **_kwargs):
            assert tools == scope
            raise PlannerFailure("planner_unavailable", "synthetic failure", retryable=True)

    failed_model = AgentRuntime(
        container, DocumentRunner(container, FakeModel()), planner=UnavailablePlanner(),
    )
    assert asyncio.run(failed_model.plan_for_worker(fallback["id"])) == 1
    state = container.agent.get(owner, fallback["id"])
    assert state["planner_mode"] == "fallback"
    assert [step["tool"] for step in state["tool_invocations"]] == scope

    with pytest.raises(CoworkerError) as unmatched:
        deterministic_plan(
            "Draft an email.", [], "en", settings, tool_allowlist=scope,
        )
    assert unmatched.value.code == "no_agent_tool"


@pytest.mark.parametrize(("retryable", "expected_non_retryable"), [(True, False), (False, True)])
def test_agent_replan_activity_preserves_planner_retryability(retryable, expected_non_retryable):
    from temporalio.exceptions import ApplicationError
    from services.coworker.agent_planning_model import PlannerFailure
    from services.coworker.worker import AgentActivities

    class FailingRuntime:
        async def replan(self, *_args, **_kwargs):
            raise PlannerFailure(
                "planner_unavailable", "The planner is temporarily unavailable.",
                retryable=retryable,
            )

    activity = AgentActivities(FailingRuntime())
    with pytest.raises(ApplicationError) as wrapped:
        asyncio.run(activity.replan({
            "run_id": "synthetic-run",
            "from_ordinal": 1,
            "reason": "result_incomplete",
        }))
    assert wrapped.value.non_retryable is expected_non_retryable


def test_agent_plan_limits_step_count(container):
    from services.coworker.agent_schemas import AgentPlanStep, AgentRunCreate
    enabled = replace(container.settings, agent_runtime_enabled=True, work_services_enabled=True)
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    owner = account(container)
    run, _ = container.agent.create(owner, AgentRunCreate(goal="Prepare several outputs.", output_language="en"), "agent-plan-limit")
    step = AgentPlanStep(tool="email.draft", arguments={
        "instruction": "Draft an email.", "notes": "", "document_ids": [], "output_language": "en",
    })
    with pytest.raises(CoworkerError) as too_many:
        container.agent.save_plan(owner, run["id"], [step] * 9)
    assert too_many.value.code == "invalid_plan"

    duplicate_run, _ = container.agent.create(
        owner,
        AgentRunCreate(goal="Prepare one email only.", output_language="en"),
        "agent-plan-duplicate-step",
    )
    with pytest.raises(CoworkerError) as duplicate:
        container.agent.save_plan(owner, duplicate_run["id"], [step, step])
    assert duplicate.value.code == "agent_duplicate_step"



def test_oauth_provider_binding_model_matches_portable_migration(container):
    from sqlalchemy import inspect
    from services.coworker.models import OAuthAttempt

    columns = {
        item["name"]: item
        for item in inspect(
            container.repository.sessions.kw["bind"]
        ).get_columns("cw_oauth_attempts")
    }
    assert set(columns) == set(OAuthAttempt.__table__.columns.keys())
    assert "provider" in columns
    assert columns["provider"]["nullable"] is False
    assert "google" in str(columns["provider"].get("default", "")).lower()


def test_agent_runtime_model_matches_migration(container):
    from sqlalchemy import inspect
    from services.coworker.models import AgentRun
    database_columns = {item["name"] for item in inspect(container.repository.sessions.kw["bind"]).get_columns("cw_agent_runs")}
    model_columns = set(AgentRun.__table__.columns.keys())
    assert database_columns == model_columns
    assert "event_sequence" in model_columns
    assert {"planner_calls", "planner_tokens", "planner_mode"} <= model_columns



def test_deterministic_agent_planner_routes_only_enabled_non_consequential_tools(container):
    from services.coworker.agent_planner import deterministic_plan
    enabled = replace(container.settings, work_services_enabled=True, artifact_services_enabled=True,
                      research_services_enabled=True, actions_enabled=True)
    plan = deterministic_plan(
        "Research competitors and prepare a presentation with a follow-up email.",
        [], "en", enabled,
    )
    assert [step.tool for step in plan] == ["research.search", "presentation.create", "email.draft"]
    assert all(step.tool not in {"email.send", "calendar.create"} for step in plan)
    assert all(step.arguments["output_language"] == "en" for step in plan)


def test_agent_outbox_is_crash_safe_and_flag_independent(container):
    from services.coworker.agent_schemas import AgentRunCreate
    from services.coworker.models import AgentOutbox
    enabled = replace(container.settings, agent_runtime_enabled=True, work_services_enabled=True)
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    owner = account(container)
    run, _ = container.agent.create(owner, AgentRunCreate(
        goal="Draft a professional email.", output_language="en"
    ), "agent-outbox-test")
    worker_id = container.repository.register_worker(
        task_queue=container.settings.task_queue,
        source_revision=container.settings.source_revision,
    )
    assert container.repository.heartbeat_worker(worker_id) == "running"
    assert container.agent.claim_outbox() == []  # No identity means no dispatch authority.
    claimed = container.agent.claim_outbox(worker_id=worker_id)
    assert claimed == [run["id"]]
    assert container.repository.complete_worker_claim(
        worker_id=worker_id, claim_type="agent_dispatch", resource_id=run["id"],
    ) is True
    with container.repository.sessions.begin() as db:
        row = db.get(AgentOutbox, run["id"])
        row.lease_until = utcnow() - timedelta(seconds=1)
    assert container.agent.claim_outbox(worker_id=worker_id) == [run["id"]]
    assert container.repository.request_worker_drain(
        task_queue=container.settings.task_queue,
        source_revision=container.settings.source_revision or "unknown",
        actor_reference="cw11-test-operator",
        reason="verify an active Agent dispatch claim blocks drain completion",
    ) == 1
    assert container.agent.claim_outbox(worker_id=worker_id) == []
    with pytest.raises(CoworkerError) as blocked:
        container.repository.mark_worker_drained(worker_id)
    assert blocked.value.code == "worker_claims_active"
    container.agent.delivered(run["id"])
    assert container.repository.complete_worker_claim(
        worker_id=worker_id, claim_type="agent_dispatch", resource_id=run["id"],
    ) is True
    container.repository.mark_worker_drained(worker_id)


def test_agent_outbox_model_matches_migration(container):
    from sqlalchemy import inspect
    from services.coworker.models import AgentOutbox
    database_columns = {item["name"] for item in inspect(container.repository.sessions.kw["bind"]).get_columns("cw_agent_outbox")}
    assert database_columns == set(AgentOutbox.__table__.columns.keys())



def test_agent_run_binds_only_owned_pending_actions(container):
    from action_samples import enable_actions, connected, action_request
    from services.coworker.agent_schemas import AgentRunCreate
    provider = enable_actions(container)
    enabled = replace(container.settings, agent_runtime_enabled=True, work_services_enabled=True)
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    container.actions.repo.settings = enabled
    owner = account(container)
    connection = connected(container.actions.repo, owner)
    action = container.actions.repo.prepare(owner, action_request(connection), "agent-bound-preview")
    run, _ = container.agent.create(owner, AgentRunCreate(
        goal="Complete the attached external action.",
        action_ids=[action["id"]],
        output_language="en",
    ), "agent-action-run")
    assert run["action_ids"] == [action["id"]]
    assert provider.sent == []

    with pytest.raises(CoworkerError) as duplicate_state:
        approved = container.actions.repo.approve(owner, action["id"], action["preview_hash"])
        container.agent.create(owner, AgentRunCreate(
            goal="Try to attach an already approved action.",
            action_ids=[approved["id"]],
            output_language="en",
        ), "agent-action-approved")
    assert duplicate_state.value.code == "action_not_awaiting_approval"


def test_agent_run_rejects_reminder_action_binding(container):
    from action_samples import enable_actions, connected, action_request
    from services.coworker.agent_schemas import AgentRunCreate

    enable_actions(container)
    enabled = replace(
        container.settings,
        agent_runtime_enabled=True,
        work_services_enabled=True,
        action_reminders_enabled=True,
    )
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    container.actions.repo.settings = enabled

    owner = account(container)
    connection = connected(container.actions.repo, owner, "calendar")
    action = container.actions.repo.prepare(
        owner,
        action_request(connection, "calendar_create_with_reminder"),
        "agent-reminder-preview",
    )

    with pytest.raises(CoworkerError) as rejected:
        container.agent.create(
            owner,
            AgentRunCreate(
                goal="Use the attached reminder action.",
                action_ids=[action["id"]],
                output_language="en",
            ),
            "agent-reminder-run",
        )
    assert rejected.value.code == "action_not_agent_selectable"


def test_agent_plan_rejects_unattached_action(container):
    from action_samples import enable_actions, connected, action_request
    from services.coworker.agent_schemas import AgentPlanStep, AgentRunCreate
    enable_actions(container)
    enabled = replace(container.settings, agent_runtime_enabled=True, work_services_enabled=True)
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    container.actions.repo.settings = enabled
    owner = account(container)
    connection = connected(container.actions.repo, owner)
    action = container.actions.repo.prepare(owner, action_request(connection), "agent-unbound-preview")
    run, _ = container.agent.create(owner, AgentRunCreate(
        goal="Create a document.",
        output_language="en",
    ), "agent-unbound-run")
    with pytest.raises(CoworkerError) as scope:
        container.agent.save_plan(owner, run["id"], [
            AgentPlanStep(tool="email.send", arguments={"action_id": action["id"]}),
        ])
    assert scope.value.code == "action_scope"


def test_agent_action_binding_model_matches_migration(container):
    from sqlalchemy import inspect
    from services.coworker.models import AgentRun
    columns = {item["name"] for item in inspect(container.repository.sessions.kw["bind"]).get_columns("cw_agent_runs")}
    assert columns == set(AgentRun.__table__.columns.keys())
    assert "action_ids" in columns



def test_agent_bound_action_cannot_dispatch_before_its_step(container):
    from action_samples import enable_actions, connected, action_request
    from services.coworker.agent_schemas import AgentPlanStep, AgentRunCreate
    provider = enable_actions(container)
    enabled = replace(container.settings, agent_runtime_enabled=True, work_services_enabled=True)
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    container.actions.repo.settings = enabled
    owner = account(container)
    connection = connected(container.actions.repo, owner)
    action = container.actions.repo.prepare(owner, action_request(connection), "agent-early-approval-preview")
    run, _ = container.agent.create(owner, AgentRunCreate(
        goal="Complete the attached external action.",
        action_ids=[action["id"]],
        output_language="en",
    ), "agent-early-approval-run")

    approved = container.actions.repo.approve(owner, action["id"], action["preview_hash"])
    assert approved["state"] == "queued"
    assert container.actions.repo.claim_outbox() == []
    assert provider.sent == []

    container.agent.save_plan(owner, run["id"], [
        AgentPlanStep(tool="email.send", arguments={"action_id": action["id"]}),
    ])
    container.agent.action_waiting(run["id"], 1, action["id"], "queued")
    worker_id = container.repository.register_worker(
        task_queue=container.settings.task_queue,
        source_revision=container.settings.source_revision,
    )
    assert container.repository.heartbeat_worker(worker_id) == "running"
    assert container.actions.repo.claim_outbox() == []
    assert container.actions.repo.claim_outbox(worker_id=worker_id) == [action["id"]]
    assert container.repository.request_worker_drain(
        task_queue=container.settings.task_queue,
        source_revision=container.settings.source_revision or "unknown",
        actor_reference="cw11-test-operator",
        reason="verify an active action dispatch claim blocks drain completion",
    ) == 1
    assert container.actions.repo.claim_outbox(worker_id=worker_id) == []
    with pytest.raises(CoworkerError) as blocked:
        container.repository.mark_worker_drained(worker_id)
    assert blocked.value.code == "worker_claims_active"
    assert container.repository.complete_worker_claim(
        worker_id=worker_id, claim_type="action_dispatch", resource_id=action["id"],
    ) is True
    container.actions.repo.delivered(action["id"])
    container.repository.mark_worker_drained(worker_id)


def test_agent_dispatch_requires_an_active_worker_claim(container):
    from services.coworker.agent_schemas import AgentRunCreate
    from services.coworker.worker import Dispatcher

    enabled = replace(
        container.settings, agent_runtime_enabled=True, work_services_enabled=True,
    )
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    owner = account(container)
    run, _ = container.agent.create(
        owner,
        AgentRunCreate(goal="Prepare a bounded work plan.", output_language="en"),
        "agent-dispatch-without-claim",
    )
    worker_id = container.repository.register_worker(
        task_queue=container.settings.task_queue,
        source_revision=container.settings.source_revision,
    )
    assert container.repository.heartbeat_worker(worker_id) == "running"

    with pytest.raises(CoworkerError) as blocked:
        asyncio.run(Dispatcher(container, object(), worker_id).dispatch_agent_run(run["id"]))
    assert blocked.value.code == "worker_claim_missing"
    container.repository.mark_worker_drained(worker_id)


def test_dispatch_batch_failure_closes_unattempted_claims(container):
    from services.coworker.models import WorkerClaim, WorkerInstance, utcnow
    from services.coworker.worker import Dispatcher
    from services.coworker.worker_claims import add_worker_claim

    worker_id = container.repository.register_worker(
        task_queue=container.settings.task_queue,
        source_revision=container.settings.source_revision,
    )
    assert container.repository.heartbeat_worker(worker_id) == "running"
    resource_ids = ["first-dispatch", "unattempted-dispatch"]
    with container.repository.sessions.begin() as db:
        worker = db.get(WorkerInstance, worker_id)
        for resource_id in resource_ids:
            add_worker_claim(
                db,
                worker=worker,
                claim_type="task_dispatch",
                resource_id=resource_id,
                now=utcnow(),
            )

    attempted = []

    async def fail_first(resource_id):
        attempted.append(resource_id)
        raise RuntimeError("simulated Temporal dispatch failure")

    async def run_batch():
        await Dispatcher(container, object(), worker_id)._dispatch_claim_batch(
            "task_dispatch", resource_ids, fail_first,
        )

    with pytest.raises(RuntimeError, match="simulated Temporal dispatch failure"):
        asyncio.run(run_batch())

    assert attempted == ["first-dispatch"]
    with container.repository.sessions() as db:
        claims = db.scalars(select(WorkerClaim).where(
            WorkerClaim.worker_id == worker_id,
        )).all()
        assert {claim.resource_id: claim.state for claim in claims} == {
            "first-dispatch": "completed",
            "unattempted-dispatch": "completed",
        }
        assert all(claim.completed_at is not None for claim in claims)
    assert container.repository.worker_inventory()[0]["active_claim_count"] == 0
    container.repository.mark_worker_drained(worker_id)


def test_worker_claim_model_matches_migration(container):
    from sqlalchemy import inspect
    from services.coworker.models import WorkerClaim

    inspector = inspect(container.repository.sessions.kw["bind"])
    database_columns = {
        item["name"] for item in inspector.get_columns("cw_worker_claims")
    }
    database_indexes = {
        item["name"] for item in inspector.get_indexes("cw_worker_claims")
    }
    model_indexes = {index.name for index in WorkerClaim.__table__.indexes}

    assert database_columns == set(WorkerClaim.__table__.columns.keys())
    assert database_indexes == model_indexes


def test_external_action_binding_model_matches_migration(container):
    from sqlalchemy import inspect
    from services.coworker.models import ExternalAction
    columns = {item["name"] for item in inspect(container.repository.sessions.kw["bind"]).get_columns("cw_external_actions")}
    assert columns == set(ExternalAction.__table__.columns.keys())
    assert {"agent_run_id", "agent_ready"} <= columns



def test_structured_memory_owned_versioned_expiring_and_hard_deletable(container):
    from services.coworker.memory_schemas import MemoryFactCreate, MemoryFactUpdate
    from services.coworker.models import AuditEvent, MemoryFact
    enabled = replace(container.settings, agent_memory_enabled=True)
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    container.memory.settings = enabled
    alice, bob = account(container), account(container, "bob")

    fact = container.memory.create(alice, MemoryFactCreate(
        namespace="preferences", key="writing.tone", value="Use concise professional English.", language="en"
    ))
    assert fact["version"] == 1
    assert container.memory.list(bob) == []
    with pytest.raises(CoworkerError):
        container.memory.update(bob, fact["id"], MemoryFactUpdate(value="tamper", language="en"))

    updated = container.memory.update(alice, fact["id"], MemoryFactUpdate(
        value="Use concise professional US English.", language="en",
        expires_at=utcnow() + timedelta(days=1),
    ))
    assert updated["version"] == 2
    from services.coworker.agent_schemas import AgentRunCreate
    container.settings = replace(enabled, agent_runtime_enabled=True)
    container.repository.settings = container.settings
    container.agent.settings = container.settings
    container.memory.settings = container.settings
    scoped_run, _ = container.agent.create(alice, AgentRunCreate(
        goal="Use my preferences.", memory_namespaces=["preferences"], output_language="en"
    ), "memory-context-scope")
    context = container.memory.context_for_run(alice, scoped_run["id"])
    assert context["facts"][0]["value"] == "Use concise professional US English."
    assert context["provenance"] == [{"id": fact["id"], "version": 2}]

    container.memory.delete(alice, fact["id"])
    assert container.memory.list(alice) == []
    assert container.memory.context_for_run(alice, scoped_run["id"]) == {"facts": [], "provenance": []}
    with container.repository.sessions() as db:
        assert db.get(MemoryFact, fact["id"]) is None
        audits = db.scalars(select(AuditEvent).where(AuditEvent.resource_id == fact["id"])).all()
        assert [event.action for event in audits] == ["memory.created", "memory.updated", "memory.deleted"]
        assert all("professional" not in event.action for event in audits)


def test_memory_create_update_disabled_but_delete_remains_available(container):
    from services.coworker.memory_schemas import MemoryFactCreate
    enabled = replace(container.settings, agent_memory_enabled=True)
    container.settings = enabled
    container.memory.settings = enabled
    owner = account(container)
    fact = container.memory.create(owner, MemoryFactCreate(
        namespace="profile", key="display_name", value="Alice", language="en"
    ))
    container.memory.settings = replace(enabled, agent_memory_enabled=False)
    with pytest.raises(CoworkerError) as disabled:
        container.memory.create(owner, MemoryFactCreate(
            namespace="profile", key="company", value="Example", language="en"
        ))
    assert disabled.value.code == "agent_memory_unavailable"
    assert container.memory.list(owner)[0]["id"] == fact["id"]
    assert container.memory.delete(owner, fact["id"])["deleted"] is True


def test_memory_model_matches_migration(container):
    from sqlalchemy import inspect
    from services.coworker.models import MemoryFact, Task
    inspector = inspect(container.repository.sessions.kw["bind"])
    memory_columns = {item["name"] for item in inspector.get_columns("cw_memory_facts")}
    task_columns = {item["name"] for item in inspector.get_columns("cw_tasks")}
    assert memory_columns == set(MemoryFact.__table__.columns.keys())
    assert task_columns == set(Task.__table__.columns.keys())
    assert {"agent_run_id", "agent_step_id"} <= task_columns


def test_agent_memory_is_ephemeral_and_receipt_keeps_only_provenance(container):
    from services.coworker.agent_runtime import AgentRuntime
    from services.coworker.agent_schemas import AgentRunCreate
    from services.coworker.memory_schemas import MemoryFactCreate

    enabled = replace(container.settings, agent_runtime_enabled=True, agent_memory_enabled=True)
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    container.memory.settings = enabled
    owner = account(container)
    fact = container.memory.create(owner, MemoryFactCreate(
        namespace="preferences", key="writing.tone", value="MEMORY-PRIVATE-SENTINEL", language="en"
    ))
    profile_fact = container.memory.create(owner, MemoryFactCreate(
        namespace="profile", key="display_name", value="PROFILE-PRIVATE-SENTINEL", language="en"
    ))

    class CaptureModel(FakeModel):
        def __init__(self):
            super().__init__()
            self.memory_seen = []
        def messages(self, task, sources):
            self.memory_seen.append(task.get("memory"))
            return super().messages(task, sources)

    run, _ = container.agent.create(owner, AgentRunCreate(
        goal="Prepare a project report.", memory_namespaces=["preferences"], output_language="en"
    ), "agent-memory-run")
    model = CaptureModel()
    runtime = AgentRuntime(container, DocumentRunner(container, model))
    assert runtime.plan(run["id"]) == 1
    asyncio.run(runtime.execute_step(run["id"], 1))
    runtime.complete(run["id"])
    saved = container.agent.get(owner, run["id"])
    assert model.memory_seen == [[{
        "namespace": "preferences", "key": "writing.tone",
        "value": "MEMORY-PRIVATE-SENTINEL", "language": "en",
    }]]
    receipt = saved["tool_invocations"][0]["receipt"]
    assert receipt["summary"]["memory"] == [{"id": fact["id"], "version": 1}]
    assert profile_fact["id"] not in json.dumps(receipt)
    assert "MEMORY-PRIVATE-SENTINEL" not in json.dumps(receipt)
    assert "PROFILE-PRIVATE-SENTINEL" not in json.dumps(receipt)

    container.memory.delete(owner, fact["id"])
    second, _ = container.agent.create(owner, AgentRunCreate(
        goal="Prepare another project report.", memory_namespaces=["preferences"], output_language="en"
    ), "agent-memory-run-2")
    second_model = CaptureModel()
    second_runtime = AgentRuntime(container, DocumentRunner(container, second_model))
    second_runtime.plan(second["id"])
    asyncio.run(second_runtime.execute_step(second["id"], 1))
    assert second_model.memory_seen == [None]


def test_standalone_task_never_receives_agent_memory(container):
    from services.coworker.memory_schemas import MemoryFactCreate
    enabled = replace(container.settings, agent_memory_enabled=True)
    container.settings = enabled
    container.repository.settings = enabled
    container.memory.settings = enabled
    owner = account(container)
    container.memory.create(owner, MemoryFactCreate(
        namespace="preferences", key="writing.tone", value="Do not leak into standalone task.", language="en"
    ))

    class CaptureModel(FakeModel):
        def __init__(self):
            super().__init__()
            self.memory_seen = []
        def messages(self, task, sources):
            self.memory_seen.append(task.get("memory"))
            return super().messages(task, sources)

    task = new_task(container, owner)
    model = CaptureModel()
    asyncio.run(DocumentRunner(container, model).run_for_test(task["id"]))
    assert model.memory_seen == [None]



def test_memory_api_is_owner_scoped_and_user_controlled(signed_client, container):
    client, headers = signed_client
    enabled = replace(container.settings, agent_memory_enabled=True)
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    container.memory.settings = enabled
    alice, bob = headers(), headers("bob")
    created = client.post("/api/v1/memory", headers=alice, json={
        "namespace": "preferences", "key": "writing.tone",
        "value": "Concise professional English.", "language": "en",
    })
    assert created.status_code == 201
    fact = created.json()
    assert client.get("/api/v1/memory", headers=bob).json()["facts"] == []
    assert client.put(f'/api/v1/memory/{fact["id"]}', headers=bob, json={
        "value": "tamper", "language": "en",
    }).status_code == 404
    assert client.delete(f'/api/v1/memory/{fact["id"]}', headers=bob).status_code == 404
    updated = client.put(f'/api/v1/memory/{fact["id"]}', headers=alice, json={
        "value": "Concise professional US English.", "language": "en",
    })
    assert updated.status_code == 200 and updated.json()["version"] == 2
    assert client.delete(f'/api/v1/memory/{fact["id"]}', headers=alice).json()["deleted"] is True
    assert client.get("/api/v1/memory", headers=alice).json()["facts"] == []


def test_agent_memory_requires_explicit_namespace_scope(container):
    from services.coworker.agent_runtime import AgentRuntime
    from services.coworker.agent_schemas import AgentRunCreate
    from services.coworker.memory_schemas import MemoryFactCreate

    enabled = replace(container.settings, agent_runtime_enabled=True, agent_memory_enabled=True)
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    container.memory.settings = enabled
    owner = account(container)
    container.memory.create(owner, MemoryFactCreate(
        namespace="preferences", key="writing.tone", value="PRIVATE-MEMORY", language="en"
    ))

    class CaptureModel(FakeModel):
        def __init__(self):
            super().__init__()
            self.memory_seen = []
        def messages(self, task, sources):
            self.memory_seen.append(task.get("memory"))
            return super().messages(task, sources)

    run, _ = container.agent.create(owner, AgentRunCreate(
        goal="Prepare a project report.", output_language="en"
    ), "agent-memory-no-scope")
    model = CaptureModel()
    runtime = AgentRuntime(container, DocumentRunner(container, model))
    runtime.plan(run["id"])
    asyncio.run(runtime.execute_step(run["id"], 1))
    assert model.memory_seen == [None]



def test_expired_memory_stays_user_visible_but_not_in_agent_context(container):
    from services.coworker.agent_schemas import AgentRunCreate
    from services.coworker.memory_schemas import MemoryFactCreate
    from services.coworker.models import MemoryFact

    enabled = replace(container.settings, agent_runtime_enabled=True, agent_memory_enabled=True)
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    container.memory.settings = enabled
    owner = account(container)
    fact = container.memory.create(owner, MemoryFactCreate(
        namespace="preferences", key="old.preference", value="Expired value", language="en"
    ))
    with container.repository.sessions.begin() as db:
        db.get(MemoryFact, fact["id"]).expires_at = utcnow() - timedelta(seconds=1)

    listed = container.memory.list(owner)
    assert listed[0]["id"] == fact["id"] and listed[0]["active"] is False
    run, _ = container.agent.create(owner, AgentRunCreate(
        goal="Prepare a project report.", memory_namespaces=["preferences"], output_language="en"
    ), "expired-memory-scope")
    assert container.memory.context_for_run(owner, run["id"]) == {"facts": [], "provenance": []}



def test_intelligent_planner_reconstructs_server_owned_arguments(container):
    from services.coworker.agent_planner import proposal_to_plan
    from services.coworker.agent_schemas import AgentPlannerProposal
    enabled = replace(container.settings, work_services_enabled=True, research_services_enabled=True)
    proposal = AgentPlannerProposal.model_validate({
        "steps": [
            {"tool": "research.search", "objective": "Compare the current market."},
            {"tool": "email.draft", "objective": "Draft a concise follow-up."},
        ]
    })
    plan = proposal_to_plan(proposal, "Private original goal", [], "bn", enabled)
    assert [step.tool for step in plan] == ["research.search", "email.draft"]
    assert plan[0].arguments["query"] == "Private original goal"
    assert plan[0].arguments["instruction"] == "Private original goal"
    assert plan[1].arguments["instruction"] == "Private original goal"
    assert "Compare the current market." not in json.dumps([step.arguments for step in plan])
    assert plan[0].arguments["document_ids"] == []
    assert plan[0].arguments["output_language"] == "bn"
    assert "recipient" not in json.dumps([step.arguments for step in plan]).lower()


def test_intelligent_planner_rejects_model_invented_tool(container):
    from services.coworker.agent_planner import proposal_to_plan
    from services.coworker.agent_schemas import AgentPlannerProposal
    proposal = AgentPlannerProposal.model_validate({
        "steps": [{"tool": "shell.run", "objective": "Run a command."}]
    })
    with pytest.raises(CoworkerError) as error:
        proposal_to_plan(proposal, "Do work", [], "en", replace(container.settings, work_services_enabled=True))
    assert error.value.code == "planner_tool_scope"


def test_agent_intelligent_plan_falls_back_deterministically(container):
    from services.coworker.agent_runtime import AgentRuntime
    from services.coworker.agent_schemas import AgentRunCreate
    from services.coworker.agent_planning_model import PlannerFailure

    enabled = replace(container.settings, agent_runtime_enabled=True, intelligent_planner_enabled=True,
                      work_services_enabled=True, max_agent_planner_calls=2, agent_planner_token_budget=16000)
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    owner = account(container)
    run, _ = container.agent.create(owner, AgentRunCreate(
        goal="Draft a professional follow-up email.", output_language="en"
    ), "intelligent-fallback")

    class FailingPlanner:
        async def propose(self, *_args, **_kwargs):
            raise PlannerFailure("planner_unavailable", "temporary", retryable=True)

    runtime = AgentRuntime(container, DocumentRunner(container, FakeModel()), planner=FailingPlanner())
    assert asyncio.run(runtime.plan_for_worker(run["id"])) == 1
    saved = container.agent.get(owner, run["id"])
    assert saved["planner_mode"] == "fallback"
    assert saved["planner_calls"] == 1
    assert saved["planner_tokens"] == 8000
    assert saved["steps"][0]["tool"] == "email.draft"


def test_agent_planner_budget_is_reserved_before_provider_call(container):
    from services.coworker.agent_schemas import AgentRunCreate
    enabled = replace(container.settings, agent_runtime_enabled=True, intelligent_planner_enabled=True,
                      work_services_enabled=True, max_agent_planner_calls=2, agent_planner_token_budget=100)
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    owner = account(container)
    run, _ = container.agent.create(owner, AgentRunCreate(goal="Create a document.", output_language="en"),
                                    "planner-budget")
    first = container.agent.reserve_planner(run["id"], 50)
    assert first["call"] == 1
    second = container.agent.reserve_planner(run["id"], 50)
    assert second["call"] == 2
    with pytest.raises(CoworkerError) as limit:
        container.agent.reserve_planner(run["id"], 1)
    assert limit.value.code == "planner_call_limit"
    saved = container.agent.get(owner, run["id"])
    assert saved["planner_tokens"] == 100


def test_agent_step_requests_replan_only_for_disabled_prepared_tool(container):
    from services.coworker.agent_runtime import AgentRuntime
    from services.coworker.agent_schemas import AgentRunCreate
    enabled = replace(container.settings, agent_runtime_enabled=True, intelligent_planner_enabled=True,
                      work_services_enabled=True)
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    owner = account(container)
    run, _ = container.agent.create(owner, AgentRunCreate(
        goal="Draft a professional email.", output_language="en"
    ), "replan-disabled-tool")
    runtime = AgentRuntime(container, DocumentRunner(container, FakeModel()))
    assert runtime.plan(run["id"]) == 1
    container.settings = replace(enabled, work_services_enabled=False)
    container.repository.settings = container.settings
    container.agent.settings = container.settings
    result = asyncio.run(runtime.execute_step(run["id"], 1))
    assert result == {"status": "replan_required"}



def test_intelligent_planner_http_contract_is_bounded_and_private(container):
    from services.coworker.agent_planning_model import DeepSeekAgentPlanner, PlannerFailure
    settings = replace(container.settings, deepseek_api_key="test-only-placeholder",
                       agent_planner_max_output_tokens=1200)
    seen = {}
    def respond(request):
        body = json.loads(request.content)
        seen.update(body)
        assert body["thinking"] == {"type": "disabled"}
        assert body["response_format"] == {"type": "json_object"}
        user = json.loads(body["messages"][1]["content"])
        assert user["available_tools"] == ["email.draft", "report.create"]
        assert "PRIVATE-MEMORY" not in request.content.decode()
        return httpx.Response(200, json={
            "choices": [{"finish_reason": "stop", "message": {"content": json.dumps({
                "steps": [{"tool": "email.draft", "objective": "Draft a concise update."}]
            })}}],
            "usage": {"total_tokens": 321},
        })
    planner = DeepSeekAgentPlanner(settings, httpx.MockTransport(respond))
    proposal, tokens, _latency = asyncio.run(planner.propose(
        "Prepare an update.", ["email.draft", "report.create"]
    ))
    assert proposal.steps[0].tool == "email.draft"
    assert tokens == 321
    assert seen["max_tokens"] == 1200

    bad = DeepSeekAgentPlanner(settings, httpx.MockTransport(lambda _: httpx.Response(200, json={
        "choices": [{"finish_reason": "stop", "message": {"content": json.dumps({
            "steps": [{"tool": "shell.run", "objective": "Run command."}]
        })}}],
        "usage": {"total_tokens": 12},
    })))
    with pytest.raises(PlannerFailure) as invalid:
        asyncio.run(bad.propose("Do work", ["email.draft"]))
    assert invalid.value.code == "invalid_planner_output"

    busy = DeepSeekAgentPlanner(settings, httpx.MockTransport(
        lambda _: httpx.Response(429, text="PRIVATE PROVIDER BODY")
    ))
    with pytest.raises(PlannerFailure) as failure:
        asyncio.run(busy.propose("Do work", ["email.draft"]))
    assert failure.value.retryable and "PRIVATE" not in failure.value.message



def test_agent_runtime_replan_replaces_unstarted_step_directly(container):
    from services.coworker.agent_runtime import AgentRuntime
    from services.coworker.agent_schemas import AgentPlannerProposal, AgentRunCreate

    enabled = replace(container.settings, agent_runtime_enabled=True, intelligent_planner_enabled=True,
                      work_services_enabled=True, max_agent_planner_calls=2,
                      agent_planner_token_budget=16000)
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    owner = account(container)
    run, _ = container.agent.create(owner, AgentRunCreate(
        goal="Draft a professional follow-up email.", output_language="en"
    ), "direct-replan")
    runtime = AgentRuntime(container, DocumentRunner(container, FakeModel()))
    assert runtime.plan(run["id"]) == 1

    class Planner:
        async def propose(self, *_args, **_kwargs):
            return AgentPlannerProposal.model_validate({
                "steps": [{"tool": "report.create", "objective": "Prepare a professional update."}]
            }), 10, 1

    changed = replace(enabled, work_services_enabled=False)
    container.settings = changed
    container.repository.settings = changed
    container.agent.settings = changed
    runtime.planner = Planner()
    assert asyncio.run(runtime.replan(run["id"], 1)) == 1
    saved = container.agent.get(owner, run["id"])
    assert [step["tool"] for step in saved["steps"]] == ["report.create"]
    assert saved["planner_mode"] == "replanned"



def test_agent_handoff_uses_nearest_completed_task_with_bounded_provenance(container):
    from coworker_samples import WorkModel
    from services.coworker.agent_runtime import AgentRuntime
    from services.coworker.agent_schemas import AgentPlanStep, AgentRunCreate

    enabled = replace(
        container.settings,
        agent_runtime_enabled=True,
        agent_handoffs_enabled=True,
        work_services_enabled=True,
        max_agent_handoff_bytes=160,
    )
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    owner = account(container)
    other = account(container, "bob")
    run, _ = container.agent.create(owner, AgentRunCreate(
        goal="Create a project document and then draft a follow-up email.",
        output_language="en",
    ), "agent-handoff-run")
    saved = container.agent.save_plan(owner, run["id"], [
        AgentPlanStep(tool="document.create", arguments={
            "instruction": run["goal"], "notes": "", "document_ids": [], "output_language": "en",
        }),
        AgentPlanStep(tool="email.draft", arguments={
            "instruction": run["goal"], "notes": "", "document_ids": [], "output_language": "en",
        }),
    ])

    class CaptureWorkModel(WorkModel):
        def __init__(self):
            super().__init__()
            self.seen = []
        def messages(self, task, sources):
            self.seen.append((task["skill_id"], [dict(source) for source in sources]))
            return super().messages(task, sources)

    model = CaptureWorkModel()
    runtime = AgentRuntime(container, DocumentRunner(container, model))
    asyncio.run(runtime.execute_step(run["id"], 1))

    handoff = container.agent.handoff_context(owner, run["id"], saved["steps"][1]["id"])
    assert len(handoff["sources"]) == 1
    assert handoff["provenance"][0]["ordinal"] == 1
    assert handoff["provenance"][0]["tool"] == "document.create"
    assert handoff["provenance"][0]["truncated"] is True
    assert handoff["sources"][0]["agent_invocation_id"] == handoff["provenance"][0]["invocation_id"]
    assert len(handoff["sources"][0]["text"].encode("utf-8")) <= 200
    assert handoff["sources"][0]["sha256"] == hashlib.sha256(handoff["sources"][0]["text"].encode()).hexdigest()
    with pytest.raises(CoworkerError) as scope:
        container.agent.handoff_context(other, run["id"], saved["steps"][1]["id"])
    assert scope.value.code == "handoff_scope"

    asyncio.run(runtime.execute_step(run["id"], 2))
    final = container.agent.get(owner, run["id"])
    assert [skill for skill, _sources in model.seen] == ["document", "email"]
    second_sources = model.seen[1][1]
    chained = [source for source in second_sources if source.get("agent_invocation_id")]
    assert len(chained) == 1
    assert chained[0]["id"].startswith("agent-step-")
    receipt = final["tool_invocations"][1]["receipt"]
    assert receipt["summary"]["handoff"] == handoff["provenance"]
    assert "The team completed 12 reviews." not in json.dumps(receipt)
    assert handoff["sources"][0]["text"] not in json.dumps(receipt)


def test_agent_handoff_is_empty_when_feature_is_disabled(container):
    from services.coworker.agent_schemas import AgentPlanStep, AgentRunCreate
    enabled = replace(container.settings, agent_runtime_enabled=True, work_services_enabled=True,
                      agent_handoffs_enabled=False)
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    owner = account(container)
    run, _ = container.agent.create(owner, AgentRunCreate(
        goal="Create two work outputs.", output_language="en"
    ), "agent-handoff-disabled")
    saved = container.agent.save_plan(owner, run["id"], [
        AgentPlanStep(tool="document.create", arguments={
            "instruction": run["goal"], "notes": "", "document_ids": [], "output_language": "en",
        }),
        AgentPlanStep(tool="email.draft", arguments={
            "instruction": run["goal"], "notes": "", "document_ids": [], "output_language": "en",
        }),
    ])
    assert container.agent.handoff_context(owner, run["id"], saved["steps"][1]["id"]) == {
        "sources": [], "provenance": [],
    }



def test_research_step_never_receives_prior_agent_handoff(container):
    from coworker_samples import WorkModel
    from research_samples import SimulatedResearch
    from services.coworker.agent_runtime import AgentRuntime
    from services.coworker.agent_schemas import AgentPlanStep, AgentRunCreate

    enabled = replace(
        container.settings,
        agent_runtime_enabled=True,
        agent_handoffs_enabled=True,
        work_services_enabled=True,
        research_services_enabled=True,
    )
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    enable_reviewed_pilot(container)
    owner = account(container)
    run, _ = container.agent.create(
        owner,
        AgentRunCreate(
            goal="Create a project document, then research the public project update.",
            output_language="en",
        ),
        "agent-handoff-research",
        tool_allowlist=["document.create", "research.search"],
    )
    container.agent.save_plan(owner, run["id"], [
        AgentPlanStep(tool="document.create", arguments={
            "instruction": run["goal"], "notes": "", "document_ids": [], "output_language": "en",
        }),
        AgentPlanStep(tool="research.search", arguments={
            "instruction": run["goal"], "notes": "", "document_ids": [], "output_language": "en",
            "query": "public project update", "time_range": "any",
        }),
    ])

    class CaptureWorkModel(WorkModel):
        def __init__(self):
            super().__init__()
            self.seen = []
        def messages(self, task, sources):
            self.seen.append((task["skill_id"], [dict(source) for source in sources]))
            return super().messages(task, sources)

    model = CaptureWorkModel()
    runtime = AgentRuntime(
        container,
        DocumentRunner(container, model, research=SimulatedResearch(enabled)),
    )
    asyncio.run(runtime.execute_step(run["id"], 1))
    asyncio.run(runtime.execute_step(run["id"], 2))
    assert [skill for skill, _sources in model.seen] == ["document", "research"]
    research_sources = model.seen[1][1]
    assert any(source["id"] == "web-1" for source in research_sources)
    assert all("agent_invocation_id" not in source for source in research_sources)
    receipt = container.agent.get(owner, run["id"])["tool_invocations"][1]["receipt"]
    assert receipt["summary"]["handoff"] == []


def test_agent_incomplete_result_requests_replan_only_after_completed_step(container):
    from coworker_samples import WorkModel
    from services.coworker.agent_runtime import AgentRuntime
    from services.coworker.agent_schemas import AgentPlanStep, AgentRunCreate

    enabled = replace(
        container.settings,
        agent_runtime_enabled=True,
        intelligent_planner_enabled=True,
        agent_outcome_replan_enabled=True,
        work_services_enabled=True,
    )
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    owner = account(container)
    run, _ = container.agent.create(owner, AgentRunCreate(
        goal="Create a project document and then draft a follow-up email.",
        output_language="en",
    ), "agent-outcome-replan")
    container.agent.save_plan(owner, run["id"], [
        AgentPlanStep(tool="document.create", arguments={
            "instruction": run["goal"], "notes": "", "document_ids": [], "output_language": "en",
        }),
        AgentPlanStep(tool="email.draft", arguments={
            "instruction": run["goal"], "notes": "", "document_ids": [], "output_language": "en",
        }),
    ])
    container.agent.set_planner_mode(run["id"], "intelligent")
    runtime = AgentRuntime(container, DocumentRunner(container, WorkModel(missing=True)))
    result = asyncio.run(runtime.execute_step(run["id"], 1))
    assert result == {"status": "outcome_replan_required"}
    saved = container.agent.get(owner, run["id"])
    assert saved["steps"][0]["state"] == "completed"
    assert saved["tool_invocations"][0]["receipt"]["summary"]["has_missing_information"] is True
    assert saved["tool_invocations"][1]["state"] == "prepared"


def test_agent_incomplete_result_does_not_replan_when_flag_is_off(container):
    from coworker_samples import WorkModel
    from services.coworker.agent_runtime import AgentRuntime
    from services.coworker.agent_schemas import AgentPlanStep, AgentRunCreate

    enabled = replace(
        container.settings,
        agent_runtime_enabled=True,
        intelligent_planner_enabled=True,
        agent_outcome_replan_enabled=False,
        work_services_enabled=True,
    )
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    owner = account(container)
    run, _ = container.agent.create(owner, AgentRunCreate(
        goal="Create a project document.", output_language="en"
    ), "agent-outcome-replan-off")
    container.agent.save_plan(owner, run["id"], [
        AgentPlanStep(tool="document.create", arguments={
            "instruction": run["goal"], "notes": "", "document_ids": [], "output_language": "en",
        }),
    ])
    container.agent.set_planner_mode(run["id"], "intelligent")
    runtime = AgentRuntime(container, DocumentRunner(container, WorkModel(missing=True)))
    assert asyncio.run(runtime.execute_step(run["id"], 1)) == {"status": "completed"}


def test_agent_multi_handoff_uses_two_prior_results_under_one_byte_budget(container):
    from coworker_samples import WorkModel
    from services.coworker.agent_runtime import AgentRuntime
    from services.coworker.agent_schemas import AgentPlanStep, AgentRunCreate

    enabled = replace(
        container.settings,
        agent_runtime_enabled=True,
        agent_handoffs_enabled=True,
        agent_multi_handoffs_enabled=True,
        work_services_enabled=True,
        max_agent_handoff_sources=2,
        max_agent_handoff_bytes=12000,
    )
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    owner = account(container)
    run, _ = container.agent.create(owner, AgentRunCreate(
        goal="Create a document, social update, and follow-up email.",
        output_language="en",
    ), "agent-multi-handoff")
    saved = container.agent.save_plan(owner, run["id"], [
        AgentPlanStep(tool="document.create", arguments={
            "instruction": run["goal"], "notes": "", "document_ids": [], "output_language": "en",
        }),
        AgentPlanStep(tool="social.draft", arguments={
            "instruction": run["goal"], "notes": "", "document_ids": [], "output_language": "en",
        }),
        AgentPlanStep(tool="email.draft", arguments={
            "instruction": run["goal"], "notes": "", "document_ids": [], "output_language": "en",
        }),
    ])
    runtime = AgentRuntime(container, DocumentRunner(container, WorkModel()))
    asyncio.run(runtime.execute_step(run["id"], 1))
    asyncio.run(runtime.execute_step(run["id"], 2))
    handoff = container.agent.handoff_context(owner, run["id"], saved["steps"][2]["id"])
    assert len(handoff["sources"]) == 2
    assert [item["ordinal"] for item in handoff["provenance"]] == [1, 2]
    assert sum(len(item["text"].encode("utf-8")) for item in handoff["sources"]) <= 12000
    assert all(item["sha256"] == hashlib.sha256(item["text"].encode()).hexdigest()
               for item in handoff["sources"])


def test_agent_multi_handoff_flag_off_preserves_nearest_prior_behavior(container):
    from coworker_samples import WorkModel
    from services.coworker.agent_runtime import AgentRuntime
    from services.coworker.agent_schemas import AgentPlanStep, AgentRunCreate

    enabled = replace(
        container.settings,
        agent_runtime_enabled=True,
        agent_handoffs_enabled=True,
        agent_multi_handoffs_enabled=False,
        work_services_enabled=True,
        max_agent_handoff_sources=2,
    )
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    owner = account(container)
    run, _ = container.agent.create(owner, AgentRunCreate(
        goal="Create three work outputs.", output_language="en"
    ), "agent-multi-handoff-off")
    saved = container.agent.save_plan(owner, run["id"], [
        AgentPlanStep(tool="document.create", arguments={
            "instruction": run["goal"], "notes": "", "document_ids": [], "output_language": "en",
        }),
        AgentPlanStep(tool="social.draft", arguments={
            "instruction": run["goal"], "notes": "", "document_ids": [], "output_language": "en",
        }),
        AgentPlanStep(tool="email.draft", arguments={
            "instruction": run["goal"], "notes": "", "document_ids": [], "output_language": "en",
        }),
    ])
    runtime = AgentRuntime(container, DocumentRunner(container, WorkModel()))
    asyncio.run(runtime.execute_step(run["id"], 1))
    asyncio.run(runtime.execute_step(run["id"], 2))
    handoff = container.agent.handoff_context(owner, run["id"], saved["steps"][2]["id"])
    assert len(handoff["sources"]) == 1
    assert handoff["provenance"][0]["ordinal"] == 2


def test_provider_capacity_settings_fail_closed():
    base = dict(
        database_url="sqlite://",
        auth_issuer=ISSUER,
        environment="development",
        storage_backend="local",
    )
    with pytest.raises(ValueError, match="Per-workspace provider concurrency"):
        Settings(
            **base,
            provider_max_concurrent_calls=1,
            provider_max_concurrent_per_workspace=2,
        ).validate()
    with pytest.raises(ValueError, match="token reserve"):
        Settings(
            **base,
            provider_max_reserved_tokens=1000,
            provider_max_reserved_tokens_per_workspace=2000,
        ).validate()
    with pytest.raises(ValueError, match="must exceed the model timeout"):
        Settings(
            **base,
            model_timeout_seconds=90,
            provider_lease_seconds=90,
        ).validate()


def test_global_provider_daily_budget_configuration_is_bounded():
    base = dict(
        database_url="sqlite://",
        auth_issuer=ISSUER,
        environment="development",
        storage_backend="local",
    )
    with pytest.raises(ValueError, match="global provider daily budget"):
        Settings(
            **base,
            daily_token_budget=1000000,
            provider_daily_token_budget=500000,
        ).validate()


def test_intelligent_planner_resolves_only_opaque_attached_action_handles(container):
    from services.coworker.agent_planner import (
        action_selection_candidates,
        intelligent_tool_names,
        proposal_to_plan,
    )
    from services.coworker.agent_schemas import AgentPlannerProposal

    settings = replace(
        container.settings,
        agent_runtime_enabled=True,
        intelligent_planner_enabled=True,
        actions_enabled=True,
        agent_action_selection_enabled=True,
        work_services_enabled=True,
    )
    actions = [
        {"id": "11111111-1111-1111-1111-111111111111", "kind": "email_send", "state": "awaiting_approval"},
        {"id": "22222222-2222-2222-2222-222222222222", "kind": "calendar_create", "state": "awaiting_approval"},
    ]
    candidates = action_selection_candidates(actions)
    assert list(candidates) == ["attached.email.1", "attached.calendar.2"]
    tools = intelligent_tool_names(settings, actions)
    encoded_tools = json.dumps(tools)
    assert "attached.email.1" in tools
    assert "attached.calendar.2" in tools
    assert actions[0]["id"] not in encoded_tools
    assert actions[1]["id"] not in encoded_tools

    # Run grants restrict ordinary tasks, never independently bound action handles.
    scoped_tools = intelligent_tool_names(
        settings, actions, tool_allowlist=["daily_plan.create"],
    )
    assert set(scoped_tools) == {
        "daily_plan.create", "attached.email.1", "attached.calendar.2",
    }

    proposal = AgentPlannerProposal.model_validate({
        "steps": [
            {"tool": "attached.calendar.2", "objective": "Use the attached calendar draft."},
        ]
    })
    plan = proposal_to_plan(
        proposal,
        "Schedule the reviewed event if appropriate.",
        [],
        "en",
        settings,
        actions,
    )
    assert len(plan) == 1
    assert plan[0].tool == "calendar.create"
    assert plan[0].arguments == {"action_id": actions[1]["id"]}
    scoped_plan = proposal_to_plan(
        proposal, "Schedule the reviewed event if appropriate.", [], "en",
        settings, actions, tool_allowlist=["daily_plan.create"],
    )
    assert scoped_plan == plan


def test_agent_action_selection_releases_only_unselected_pending_drafts(container):
    from action_samples import enable_actions, connected, action_request
    from services.coworker.agent_runtime import AgentRuntime
    from services.coworker.agent_schemas import AgentPlannerProposal, AgentRunCreate
    from services.coworker.models import ExternalAction

    provider = enable_actions(container)
    enabled = replace(
        container.settings,
        agent_runtime_enabled=True,
        intelligent_planner_enabled=True,
        actions_enabled=True,
        agent_action_selection_enabled=True,
        work_services_enabled=True,
        max_agent_planner_calls=2,
        agent_planner_token_budget=16000,
    )
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    container.actions.repo.settings = enabled

    owner = account(container)
    connection = connected(container.actions.repo, owner)
    first = container.actions.repo.prepare(
        owner,
        action_request(connection),
        "agent-select-first",
    )
    second = container.actions.repo.prepare(
        owner,
        action_request(connection),
        "agent-select-second",
    )
    run, _ = container.agent.create(
        owner,
        AgentRunCreate(
            goal="Use only the second attached email action draft.",
            action_ids=[first["id"], second["id"]],
            output_language="en",
        ),
        "agent-select-run",
        tool_allowlist=["daily_plan.create"],
    )

    class SelectingPlanner:
        def __init__(self):
            self.tools = None

        async def propose(self, _goal, tools, **_kwargs):
            self.tools = list(tools)
            # Replanning renumbers surviving handles after unselected drafts release.
            selected = "attached.email.2" if "attached.email.2" in tools else "attached.email.1"
            return AgentPlannerProposal.model_validate({
                "steps": [
                    {
                        "tool": selected,
                        "objective": "Use the second attached draft.",
                    },
                ]
            }), 25, 1

    planner = SelectingPlanner()
    runtime = AgentRuntime(
        container,
        DocumentRunner(container, FakeModel()),
        planner=planner,
    )
    assert asyncio.run(runtime.plan_for_worker(run["id"])) == 1

    assert planner.tools is not None
    encoded_tools = json.dumps(planner.tools)
    assert "attached.email.1" in planner.tools
    assert "attached.email.2" in planner.tools
    assert set(planner.tools) == {
        "daily_plan.create", "attached.email.1", "attached.email.2",
    }
    assert first["id"] not in encoded_tools
    assert second["id"] not in encoded_tools

    saved = container.agent.get(owner, run["id"])
    assert saved["action_ids"] == [second["id"]]
    assert saved["steps"][0]["tool"] == "email.send"
    assert asyncio.run(runtime.replan(run["id"], 1)) == 1
    assert "attached.email.1" in planner.tools
    assert "attached.email.2" not in planner.tools
    replanned = container.agent.get(owner, run["id"])
    assert replanned["action_ids"] == [second["id"]]
    assert replanned["steps"][0]["tool"] == "email.send"
    assert provider.sent == []

    with container.repository.sessions() as db:
        first_row = db.get(ExternalAction, first["id"])
        second_row = db.get(ExternalAction, second["id"])
        assert first_row.state == "awaiting_approval"
        assert first_row.agent_run_id is None
        assert first_row.agent_ready is False
        assert second_row.state == "awaiting_approval"
        assert second_row.agent_run_id == run["id"]
        assert second_row.agent_ready is False


def test_agent_action_selection_never_model_selects_already_approved_action(container):
    from services.coworker.agent_planner import intelligent_tool_names, proposal_to_plan
    from services.coworker.agent_schemas import AgentPlannerProposal

    settings = replace(
        container.settings,
        agent_runtime_enabled=True,
        intelligent_planner_enabled=True,
        actions_enabled=True,
        agent_action_selection_enabled=True,
        work_services_enabled=True,
    )
    actions = [
        {"id": "11111111-1111-1111-1111-111111111111", "kind": "email_send", "state": "queued"},
    ]
    tools = intelligent_tool_names(settings, actions)
    assert all(not name.startswith("attached.") for name in tools)

    proposal = AgentPlannerProposal.model_validate({
        "steps": [{"tool": "email.draft", "objective": "Draft supporting context."}]
    })
    plan = proposal_to_plan(
        proposal,
        "Prepare context and continue the already-approved action.",
        [],
        "en",
        settings,
        actions,
    )
    assert [step.tool for step in plan] == ["email.draft", "email.send"]
    assert plan[1].arguments == {"action_id": actions[0]["id"]}


def test_intelligent_planner_action_proposal_is_persisted_but_not_executable(container):
    from action_samples import enable_actions
    from services.coworker.agent_runtime import AgentRuntime
    from services.coworker.agent_schemas import AgentPlannerProposal, AgentRunCreate
    from services.coworker.models import ActionProposal, ExternalAction

    provider = enable_actions(container)
    enabled = replace(
        container.settings,
        agent_runtime_enabled=True,
        intelligent_planner_enabled=True,
        agent_action_proposals_enabled=True,
        work_services_enabled=True,
        max_agent_planner_calls=2,
        agent_planner_token_budget=16000,
    )
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    container.actions.repo.settings = enabled

    owner = account(container)
    run, _ = container.agent.create(
        owner,
        AgentRunCreate(
            goal=(
                "Draft an update and suggest an email to recipient@example.org "
                "with subject Project update and body The project is ready."
            ),
            output_language="en",
        ),
        "agent-proposal-run",
    )

    class ProposalPlanner:
        async def propose(self, _goal, _tools, **kwargs):
            assert kwargs["allow_action_proposals"] is True
            return AgentPlannerProposal.model_validate({
                "steps": [{
                    "tool": "email.draft",
                    "objective": "Draft the supporting email content.",
                }],
                "action_proposals": [{
                    "payload": {
                        "kind": "email_send",
                        "to": ["recipient@example.org"],
                        "cc": [],
                        "bcc": [],
                        "subject": "Project update",
                        "body": "The project is ready.",
                    },
                    "rationale": "The user explicitly requested this email.",
                }],
            }), 40, 1

    runtime = AgentRuntime(
        container,
        DocumentRunner(container, FakeModel()),
        planner=ProposalPlanner(),
    )
    assert asyncio.run(runtime.plan_for_worker(run["id"])) == 1

    saved = container.agent.get(owner, run["id"])
    assert len(saved["action_proposals"]) == 1
    proposal = saved["action_proposals"][0]
    assert proposal["kind"] == "email_send"
    assert proposal["state"] == "suggested"
    assert proposal["promoted_action_id"] is None
    assert len(proposal["proposal_hash"]) == 64

    with container.repository.sessions() as db:
        assert db.get(ActionProposal, proposal["id"]) is not None
        assert db.scalars(select(ExternalAction).where(
            ExternalAction.owner_id == owner
        )).all() == []
    assert provider.requests == []


def test_agent_action_proposal_flag_dependencies_fail_closed(container):
    base = replace(
        container.settings,
        actions_enabled=False,
        agent_runtime_enabled=True,
        intelligent_planner_enabled=True,
        agent_action_proposals_enabled=True,
    )
    with pytest.raises(ValueError, match="SHUDDHO_ACTIONS_ENABLED"):
        base.validate()


def test_runtime_manifest_is_authenticated_sanitized_and_uncached(
    signed_client,
    container,
    monkeypatch,
):
    client, headers = signed_client
    monkeypatch.setenv("SHUDDHO_COWORKER_ENABLED", "true")
    admitted = client.get(
        "/api/v1/me",
        headers=headers(),
    ).json()["account_id"]
    current = replace(
        container.settings,
        source_revision="a" * 40,
        work_services_enabled=True,
        artifact_services_enabled=True,
        agent_runtime_enabled=True,
        intelligent_planner_enabled=True,
        actions_enabled=True,
        agent_action_proposals_enabled=True,
        cohort_enforced=True,
        cohort_account_ids=frozenset({admitted}),
    )
    container.settings = current
    container.repository.settings = current
    container.agent.settings = current
    container.actions.repo.settings = current

    assert client.get("/api/v1/runtime-manifest").status_code == 401
    assert client.get(
        "/api/v1/runtime-manifest",
        headers=headers("bob"),
    ).status_code == 403
    response = client.get(
        "/api/v1/runtime-manifest",
        headers=headers(),
    )
    assert response.status_code == 200
    assert response.headers["Cache-Control"] == "no-store"
    value = response.json()
    assert value["schema_version"] == 2
    assert value["source_revision"] == "a" * 40
    assert value["capabilities"]["coworker"] is True
    assert value["capabilities"]["action_proposals"] is True
    assert value["capabilities"]["agent_browser_read"] is False
    assert value["capabilities"]["action_selection"] is False
    assert value["capabilities"]["action_reminders"] is False
    assert value["capabilities"]["action_document_sharing"] is False
    assert value["action_providers"] == ["google"]
    assert value["cohort"] == {
        "enforced": True,
        "configured_members": 1,
        "max_users": current.cohort_max_users,
    }
    encoded = json.dumps(value)
    for secret in (
        current.database_url,
        current.connector_encryption_key,
        current.deepseek_api_key,
        current.temporal_api_key,
    ):
        if secret:
            assert secret not in encoded

    # Exercise the enabled case in the API module that actually owns the
    # signed_client fixture; the isolated browser gate has no such fixture.
    activated_settings = replace(
        current, browser_enabled=True, agent_browser_read_enabled=True,
    )
    container.settings = activated_settings
    container.repository.settings = activated_settings
    container.agent.settings = activated_settings
    enabled_response = client.get("/api/v1/runtime-manifest", headers=headers())
    assert enabled_response.status_code == 200
    assert enabled_response.headers["Cache-Control"] == "no-store"
    assert enabled_response.json()["capabilities"]["browser"] is True
    assert enabled_response.json()["capabilities"]["agent_browser_read"] is True
    assert client.get("/api/v1/runtime-manifest").status_code == 401


def test_transaction_authority_manifest_is_authenticated_scoped_and_exact(
    signed_client,
    container,
    monkeypatch,
):
    client, headers = signed_client
    monkeypatch.setenv("SHUDDHO_COWORKER_ENABLED", "true")
    admitted = client.get("/api/v1/me", headers=headers()).json()["account_id"]
    current = replace(
        container.settings,
        source_revision="b" * 40,
        cohort_enforced=True,
        cohort_account_ids=frozenset({admitted}),
        personal_transactions_enabled=True,
        transaction_operations=frozenset({"google:negotiation_commitment_email"}),
    )
    container.settings = current
    container.repository.settings = current
    container.agent.settings = current
    container.actions.repo.settings = current

    assert client.get("/api/v1/transaction-authority-manifest").status_code == 401
    assert client.get(
        "/api/v1/transaction-authority-manifest",
        headers=headers("bob"),
    ).status_code == 403

    response = client.get(
        "/api/v1/transaction-authority-manifest",
        headers=headers(),
    )
    assert response.status_code == 200
    assert response.headers["Cache-Control"] == "no-store"
    assert response.json() == {
        "schema_version": 2,
        "source_revision": "b" * 40,
        "personal_transactions_enabled": True,
        "restaurant_reservations_enabled": False,
        "shopping_checkout_enabled": False,
        "travel_booking_enabled": False,
        "operations": ["google:negotiation_commitment_email"],
    }
    encoded = response.text
    for secret in (
        current.database_url,
        current.connector_encryption_key,
        current.deepseek_api_key,
        current.temporal_api_key,
    ):
        if secret:
            assert secret not in encoded


def test_core_agent_phase2_exposes_explicit_execution_lifecycle(container):
    from coworker_samples import WorkModel
    from services.coworker.agent_runtime import AgentRuntime
    from services.coworker.agent_schemas import AgentPlanStep, AgentRunCreate

    enabled = replace(
        container.settings,
        agent_runtime_enabled=True,
        work_services_enabled=True,
        max_active_agent_runs=10,
    )
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled

    owner = account(container)
    run, _ = container.agent.create(
        owner,
        AgentRunCreate(
            goal="Create a short project document.",
            output_language="en",
        ),
        "core-phase2-lifecycle",
    )
    planned = container.agent.save_plan(
        owner,
        run["id"],
        [
            AgentPlanStep(
                tool="document.create",
                arguments={
                    "instruction": "Create a short project document.",
                    "notes": "",
                    "document_ids": [],
                    "output_language": "en",
                },
            )
        ],
    )

    assert len(planned["plan"]) == 1
    assert planned["plan"][0]["step_id"] == planned["steps"][0]["id"]
    assert planned["plan"][0]["invocation_id"] == planned["tool_invocations"][0]["id"]
    assert planned["current_step"]["ordinal"] == 1
    assert planned["completed_steps"] == []
    assert planned["tool_results"] == []
    assert planned["pending_approval"] is None
    assert planned["retry_state"]["tool_steps_used"] == 1
    assert planned["retry_state"]["tool_steps_remaining"] == 7
    assert planned["final_state"] is None
    assert planned["evidence"]["verified_receipt_count"] == 0

    runtime = AgentRuntime(container, DocumentRunner(container, WorkModel()))
    assert asyncio.run(runtime.execute_step(run["id"], 1)) == {"status": "completed"}
    runtime.complete(run["id"])

    completed = container.agent.get(owner, run["id"])
    assert completed["state"] == "completed"
    assert completed["final_state"] == "completed"
    assert completed["current_step"] is None
    assert [item["ordinal"] for item in completed["completed_steps"]] == [1]
    assert len(completed["tool_results"]) == 1
    assert completed["tool_results"][0]["invocation_id"] == completed["tool_invocations"][0]["id"]
    assert completed["evidence"]["verified_receipt_count"] == 1


def test_core_agent_phase2_reports_partial_completion_when_tool_needs_input(container):
    from coworker_samples import WorkModel
    from services.coworker.agent_runtime import AgentRuntime
    from services.coworker.agent_schemas import AgentPlanStep, AgentRunCreate

    enabled = replace(
        container.settings,
        agent_runtime_enabled=True,
        work_services_enabled=True,
        max_active_agent_runs=10,
    )
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled

    owner = account(container)
    run, _ = container.agent.create(
        owner,
        AgentRunCreate(
            goal="Create a document from incomplete details.",
            output_language="en",
        ),
        "core-phase2-partial",
    )
    container.agent.save_plan(
        owner,
        run["id"],
        [
            AgentPlanStep(
                tool="document.create",
                arguments={
                    "instruction": "Create a document from incomplete details.",
                    "notes": "",
                    "document_ids": [],
                    "output_language": "en",
                },
            )
        ],
    )
    runtime = AgentRuntime(container, DocumentRunner(container, WorkModel(missing=True)))
    assert asyncio.run(runtime.execute_step(run["id"], 1)) == {"status": "completed"}
    runtime.complete(run["id"])

    result = container.agent.get(owner, run["id"])
    assert result["state"] == "completed"
    assert result["final_state"] == "partially_completed"
    assert result["tool_results"][0]["summary"]["has_missing_information"] is True


def test_core_agent_phase2_deadline_reconciliation_is_explicit(container):
    from services.coworker.agent_schemas import AgentPlanStep, AgentRunCreate
    from services.coworker.models import AgentRun

    enabled = replace(
        container.settings,
        agent_runtime_enabled=True,
        work_services_enabled=True,
        max_active_agent_runs=10,
    )
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled

    owner = account(container)
    run, _ = container.agent.create(
        owner,
        AgentRunCreate(goal="Create a project document.", output_language="en"),
        "core-phase2-deadline",
    )
    container.agent.save_plan(
        owner,
        run["id"],
        [
            AgentPlanStep(
                tool="document.create",
                arguments={
                    "instruction": "Create a project document.",
                    "notes": "",
                    "document_ids": [],
                    "output_language": "en",
                },
            )
        ],
    )
    with container.repository.sessions.begin() as db:
        db.get(AgentRun, run["id"]).deadline_at = utcnow() - timedelta(seconds=1)

    assert container.agent.expire_due_runs() == [run["id"]]
    expired = container.agent.get(owner, run["id"])
    assert expired["state"] == "failed"
    assert expired["error_code"] == "agent_deadline"
    assert expired["final_state"] == "failed"
    assert expired["steps"][0]["state"] == "failed"
    with pytest.raises(CoworkerError) as deadline:
        container.agent.assert_within_deadline(run["id"])
    assert deadline.value.code == "agent_deadline"


def test_core_agent_phase2_runtime_v3_suppresses_identical_steps(container):
    from services.coworker.agent_schemas import AgentRunCreate

    enabled = replace(
        container.settings,
        agent_runtime_enabled=True,
        intelligent_planner_enabled=True,
        agent_runtime_v3_enabled=True,
        work_services_enabled=True,
        max_active_agent_runs=10,
    )
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled

    owner = account(container)
    run, _ = container.agent.create(
        owner,
        AgentRunCreate(goal="Draft a project update email.", output_language="en"),
        "core-phase2-loop",
    )
    first = container.agent.append_v3_step(
        owner,
        run["id"],
        "email.draft",
        "Draft a project update email.",
    )
    assert first == 1

    with pytest.raises(CoworkerError) as repeated:
        container.agent.append_v3_step(
            owner,
            run["id"],
            "email.draft",
            "Draft a project update email.",
        )
    assert repeated.value.code == "agent_loop_detected"
    saved = container.agent.get(owner, run["id"])
    assert len(saved["plan"]) == 1
    assert saved["retry_state"]["tool_steps_used"] == 1


def test_core_agent_phase3_checkpoint_model_matches_migration(container):
    from sqlalchemy import inspect
    from services.coworker.models import AgentCheckpoint

    columns = {
        item["name"]
        for item in inspect(
            container.repository.sessions.kw["bind"]
        ).get_columns("cw_agent_checkpoints")
    }
    assert columns == set(AgentCheckpoint.__table__.columns.keys())


def test_core_agent_phase3_resumes_nonconsequential_step_from_durable_checkpoint(container):
    from coworker_samples import WorkModel
    from services.coworker.agent_repository import AgentRepository
    from services.coworker.agent_runtime import AgentRuntime
    from services.coworker.agent_schemas import AgentPlanStep, AgentRunCreate

    enabled = replace(
        container.settings,
        agent_runtime_enabled=True,
        work_services_enabled=True,
        max_active_agent_runs=10,
    )
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled

    owner = account(container)
    run, _ = container.agent.create(
        owner,
        AgentRunCreate(
            goal="Create a short project document.",
            output_language="en",
        ),
        "core-phase3-checkpoint-resume",
    )
    planned = container.agent.save_plan(
        owner,
        run["id"],
        [
            AgentPlanStep(
                tool="document.create",
                arguments={
                    "instruction": "Create a short project document.",
                    "notes": "",
                    "document_ids": [],
                    "output_language": "en",
                },
            )
        ],
    )
    assert planned["resume_from"] == "plan_saved"
    assert [item["kind"] for item in planned["checkpoints"]] == [
        "run_created",
        "plan_saved",
    ]

    # Simulate an API/process restart by reconstructing the repository over the
    # same durable database before any tool call.
    restarted_repo = AgentRepository(container.repository.sessions, enabled)
    recovered = restarted_repo.get(owner, run["id"])
    assert recovered["resume_from"] == "plan_saved"
    assert recovered["plan"][0]["invocation_id"] == planned["plan"][0]["invocation_id"]

    first_runtime = AgentRuntime(container, DocumentRunner(container, WorkModel()))
    assert asyncio.run(first_runtime.execute_step(run["id"], 1)) == {
        "status": "completed"
    }

    # Simulate a lost activity acknowledgement + worker replacement. The new
    # runtime must observe the completed durable invocation and must not create
    # another child task.
    second_runtime = AgentRuntime(container, DocumentRunner(container, WorkModel()))
    assert asyncio.run(second_runtime.execute_step(run["id"], 1)) == {
        "status": "completed"
    }

    with container.repository.sessions() as db:
        tasks = list(db.scalars(select(Task).where(
            Task.agent_run_id == run["id"],
        )).all())
    assert len(tasks) == 1
    assert tasks[0].idempotency_key == f'agent:{run["id"]}:1'

    recovered = restarted_repo.get(owner, run["id"])
    kinds = [item["kind"] for item in recovered["checkpoints"]]
    assert kinds.count("tool_call_started") == 1
    assert kinds.count("tool_resource_linked") == 1
    assert kinds.count("tool_observed") == 1
    assert recovered["resume_from"] == "tool_observed"

    second_runtime.complete(run["id"])
    completed = restarted_repo.get(owner, run["id"])
    assert completed["final_state"] == "completed"
    assert completed["resume_from"] == "complete"
    assert [item["kind"] for item in completed["checkpoints"]].count("complete") == 1


def test_core_agent_phase3_consequential_replay_never_mutates_provider_twice(container):
    from action_samples import action_request, connected, enable_actions
    from coworker_samples import WorkModel
    from services.coworker.agent_runtime import AgentRuntime
    from services.coworker.agent_schemas import AgentPlanStep, AgentRunCreate

    provider = enable_actions(container)
    enabled = replace(
        container.settings,
        agent_runtime_enabled=True,
        work_services_enabled=True,
        actions_enabled=True,
        max_active_agent_runs=10,
    )
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    container.actions.repo.settings = enabled

    owner = account(container)
    connection = connected(container.actions.repo, owner, "email")
    action = container.actions.repo.prepare(
        owner,
        action_request(connection, "email_send"),
        "core-phase3-action-preview",
    )
    run, _ = container.agent.create(
        owner,
        AgentRunCreate(
            goal="Send the attached approved email action.",
            action_ids=[action["id"]],
            output_language="en",
        ),
        "core-phase3-action-run",
    )
    container.agent.save_plan(
        owner,
        run["id"],
        [
            AgentPlanStep(
                tool="email.send",
                arguments={"action_id": action["id"]},
            )
        ],
    )

    runtime = AgentRuntime(container, DocumentRunner(container, WorkModel()))

    # Crash/restart before approval: repeated execution only restores the wait.
    assert asyncio.run(runtime.execute_step(run["id"], 1)) == {
        "status": "awaiting_approval"
    }
    restarted = AgentRuntime(container, DocumentRunner(container, WorkModel()))
    assert asyncio.run(restarted.execute_step(run["id"], 1)) == {
        "status": "awaiting_approval"
    }
    assert provider.sent == []

    waiting = container.agent.get(owner, run["id"])
    assert [item["kind"] for item in waiting["checkpoints"]].count(
        "wait_for_approval"
    ) == 1

    approved = container.actions.repo.approve(
        owner,
        action["id"],
        action["preview_hash"],
    )
    assert approved["state"] == "queued"

    # Crash/restart after approval but before provider mutation: the action
    # remains the same durable action ID and no second preview/approval exists.
    assert asyncio.run(restarted.execute_step(run["id"], 1)) == {
        "status": "executing"
    }
    after_approval = container.agent.get(owner, run["id"])
    assert any(
        item["kind"] == "action_approved"
        and item["resource_id"] == action["id"]
        for item in after_approval["checkpoints"]
    )

    # Simulate a provider mutation whose response is lost. The execution claim
    # is committed before the provider call, so replay may reconcile but may
    # never issue another send.
    provider.lose_reply = True
    asyncio.run(container.actions.execute(action["id"]))
    assert len(provider.sent) == 1
    first_result = container.actions.repo.get(owner, action["id"])
    assert first_result["state"] == "outcome_unknown"

    # Duplicate Temporal/activity delivery after the uncertain mutation.
    asyncio.run(container.actions.execute(action["id"]))
    assert len(provider.sent) == 1
    assert container.actions.repo.claim_execution(action["id"]) is None

    recovered = container.agent.get(owner, run["id"])
    kinds = [item["kind"] for item in recovered["checkpoints"]]
    assert kinds.count("action_execution_claimed") == 1
    assert kinds.count("action_outcome_unknown") == 1
    claimed = next(
        item for item in recovered["checkpoints"]
        if item["kind"] == "action_execution_claimed"
    )
    assert claimed["evidence"]["blind_retry_allowed"] is False
    assert claimed["evidence"]["recovery_policy"] == "reconcile_only"


def test_core_agent_phase3_verified_action_checkpoints_and_completes(container):
    from action_samples import action_request, connected, enable_actions
    from coworker_samples import WorkModel
    from services.coworker.agent_runtime import AgentRuntime
    from services.coworker.agent_schemas import AgentPlanStep, AgentRunCreate

    provider = enable_actions(container)
    enabled = replace(
        container.settings,
        agent_runtime_enabled=True,
        work_services_enabled=True,
        actions_enabled=True,
        max_active_agent_runs=10,
    )
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    container.actions.repo.settings = enabled

    owner = account(container)
    connection = connected(container.actions.repo, owner, "email")
    action = container.actions.repo.prepare(
        owner,
        action_request(connection, "email_send"),
        "core-phase3-success-preview",
    )
    run, _ = container.agent.create(
        owner,
        AgentRunCreate(
            goal="Send the attached approved email action.",
            action_ids=[action["id"]],
            output_language="en",
        ),
        "core-phase3-success-run",
    )
    container.agent.save_plan(
        owner,
        run["id"],
        [AgentPlanStep(tool="email.send", arguments={"action_id": action["id"]})],
    )
    runtime = AgentRuntime(container, DocumentRunner(container, WorkModel()))
    assert asyncio.run(runtime.execute_step(run["id"], 1)) == {
        "status": "awaiting_approval"
    }

    container.actions.repo.approve(owner, action["id"], action["preview_hash"])
    assert asyncio.run(runtime.execute_step(run["id"], 1)) == {
        "status": "executing"
    }
    asyncio.run(container.actions.execute(action["id"]))
    assert len(provider.sent) == 1

    # A worker replay after the provider receipt observes the existing action
    # and records the Agent tool receipt; it never calls the provider again.
    assert asyncio.run(
        AgentRuntime(container, DocumentRunner(container, WorkModel())).execute_step(
            run["id"], 1
        )
    ) == {"status": "completed"}
    assert len(provider.sent) == 1
    runtime.complete(run["id"])

    result = container.agent.get(owner, run["id"])
    assert result["final_state"] == "completed"
    kinds = [item["kind"] for item in result["checkpoints"]]
    assert kinds.count("action_prepared") == 1
    assert kinds.count("wait_for_approval") == 1
    assert kinds.count("action_approved") == 1
    assert kinds.count("action_execution_claimed") == 1
    assert kinds.count("action_result_verified") == 1
    assert kinds.count("verify_result") == 1
    assert kinds.count("complete") == 1
    assert result["resume_from"] == "complete"


def test_late_v3_planner_response_cannot_record_decision_or_memory_after_cancel(container):
    from services.coworker.agent_runtime import AgentRuntime
    from services.coworker.agent_schemas import AgentRunCreate, AgentV3Decision
    from services.coworker.models import AgentDecision, MemoryProposal, ToolInvocation

    enabled = replace(
        container.settings,
        agent_runtime_enabled=True,
        agent_runtime_v3_enabled=True,
        intelligent_planner_enabled=True,
        work_services_enabled=True,
        artifact_services_enabled=True,
        agent_memory_enabled=True,
        context_retrieval_enabled=True,
        max_active_agent_runs=10,
        max_agent_v3_planner_calls=2,
        agent_v3_planner_token_budget=16000,
    )
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    container.context.settings = enabled
    container.memory.settings = enabled
    owner = account(container)
    run, _ = container.agent.create(
        owner,
        AgentRunCreate(goal="Create a concise project update.", output_language="en"),
        "cw02-late-v3-cancel",
    )
    assert run["runtime_version"] == 3
    decision = AgentV3Decision.model_validate({
        "decision": "next_step",
        "tool": "email.draft",
        "objective": "Draft a concise project update.",
        "memory_proposal": {
            "namespace": "preferences",
            "key": "cancelled_run_update_style",
            "value": "Prefer concise project updates.",
            "language": "en",
            "source_ids": ["goal"],
        },
    })

    async def exercise():
        entered = asyncio.Event()
        release = asyncio.Event()

        class LatePlanner:
            async def decide(self, *_args, **_kwargs):
                entered.set()
                await release.wait()
                return decision, 24, 1, {
                    "model": "test-model",
                    "prompt_sha256": "a" * 64,
                    "tool_schema_sha256": "b" * 64,
                }

        runtime = AgentRuntime(container, DocumentRunner(container, FakeModel()), planner=LatePlanner())
        pending = asyncio.create_task(runtime.decide_v3(run["id"]))
        await asyncio.wait_for(entered.wait(), timeout=2)
        cancelled = container.agent.cancel(owner, run["id"])
        assert cancelled["state"] == "cancelled"
        release.set()
        with pytest.raises(CoworkerError) as failure:
            await pending
        assert failure.value.code == "agent_cancelled"

    asyncio.run(exercise())
    saved = container.agent.get(owner, run["id"])
    assert saved["state"] == "cancelled"
    with container.repository.sessions() as db:
        assert db.scalars(select(AgentDecision).where(
            AgentDecision.run_id == run["id"],
        )).all() == []
        assert db.scalars(select(MemoryProposal).where(
            MemoryProposal.run_id == run["id"],
        )).all() == []
        assert db.scalars(select(ToolInvocation).where(
            ToolInvocation.run_id == run["id"],
        )).all() == []


def test_cancelled_run_cannot_accept_late_memory_proposal(container):
    from services.coworker.agent_schemas import AgentMemoryProposal, AgentRunCreate
    from services.coworker.models import MemoryProposal

    enabled = replace(
        container.settings,
        agent_runtime_enabled=True,
        agent_runtime_v3_enabled=True,
        intelligent_planner_enabled=True,
        work_services_enabled=True,
        agent_memory_enabled=True,
        context_retrieval_enabled=True,
    )
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    container.context.settings = enabled
    container.memory.settings = enabled
    owner = account(container)
    run, _ = container.agent.create(
        owner,
        AgentRunCreate(goal="Use concise weekly updates.", output_language="en"),
        "cw02-cancelled-memory-proposal",
    )
    _context, source_map = container.context.planner_context(owner, run["id"])
    assert "goal" in source_map
    container.agent.cancel(owner, run["id"])
    proposal = AgentMemoryProposal(
        namespace="preferences",
        key="cancelled_run_preference",
        value="Use concise weekly updates.",
        language="en",
        source_ids=["goal"],
    )
    with pytest.raises(CoworkerError) as failure:
        container.memory.propose_from_agent(owner, run["id"], proposal, source_map)
    assert failure.value.code == "agent_cancelled"
    assert container.memory.proposals(owner) == []
    with container.repository.sessions() as db:
        assert db.scalars(select(MemoryProposal).where(
            MemoryProposal.run_id == run["id"],
        )).all() == []


def test_recorded_v3_decision_replay_remains_idempotent_after_cancel(container):
    from services.coworker.agent_schemas import AgentRunCreate, AgentV3Decision
    from services.coworker.models import AgentDecision

    enabled = replace(
        container.settings,
        agent_runtime_enabled=True,
        agent_runtime_v3_enabled=True,
        intelligent_planner_enabled=True,
        work_services_enabled=True,
        agent_memory_enabled=True,
        context_retrieval_enabled=True,
        max_active_agent_runs=10,
        max_agent_v3_planner_calls=2,
        agent_v3_planner_token_budget=16000,
    )
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    container.context.settings = enabled
    container.memory.settings = enabled
    owner = account(container)
    run, _ = container.agent.create(
        owner,
        AgentRunCreate(goal="Draft a concise project update.", output_language="en"),
        "cw02-v3-decision-replay-after-cancel",
    )
    assert run["runtime_version"] == 3
    decision = AgentV3Decision.model_validate({
        "decision": "next_step",
        "tool": "email.draft",
        "objective": "Draft a concise project update.",
    })
    record = {
        "planner_call": 1,
        "model": "test-model",
        "prompt_sha256": "a" * 64,
        "tool_schema_sha256": "b" * 64,
        "observation_count": 0,
        "total_tokens": 24,
        "cost_microusd": None,
    }

    container.agent.record_v3_decision(run["id"], decision, **record)
    cancelled = container.agent.cancel(owner, run["id"])
    assert cancelled["state"] == "cancelled"

    # An identical planner-call replay is a no-op even after the run is terminal.
    container.agent.record_v3_decision(run["id"], decision, **record)
    with container.repository.sessions() as db:
        rows = db.scalars(select(AgentDecision).where(
            AgentDecision.run_id == run["id"],
        )).all()
    assert len(rows) == 1


def test_terminal_v3_run_rejects_late_decision_and_memory_proposal(container):
    from services.coworker.agent_schemas import AgentMemoryProposal, AgentRunCreate, AgentV3Decision
    from services.coworker.models import AgentDecision, MemoryProposal

    enabled = replace(
        container.settings,
        agent_runtime_enabled=True,
        agent_runtime_v3_enabled=True,
        intelligent_planner_enabled=True,
        work_services_enabled=True,
        artifact_services_enabled=True,
        agent_memory_enabled=True,
        context_retrieval_enabled=True,
        max_active_agent_runs=10,
        max_agent_v3_planner_calls=2,
        agent_v3_planner_token_budget=16000,
    )
    container.settings = enabled
    container.repository.settings = enabled
    container.agent.settings = enabled
    container.context.settings = enabled
    container.memory.settings = enabled
    owner = account(container)
    run, _ = container.agent.create(
        owner,
        AgentRunCreate(goal="Prepare a concise project update.", output_language="en"),
        "cw02-v3-terminal-late-write",
    )
    _context, source_map = container.context.planner_context(owner, run["id"])
    container.agent.v3_terminal(run["id"], "blocked", "No registered tool is available.")

    decision = AgentV3Decision.model_validate({
        "decision": "next_step",
        "tool": "email.draft",
        "objective": "Draft a concise project update.",
    })
    with pytest.raises(CoworkerError) as failure:
        container.agent.record_v3_decision(
            run["id"], decision,
            planner_call=1,
            model="test-model",
            prompt_sha256="a" * 64,
            tool_schema_sha256="b" * 64,
            observation_count=0,
            total_tokens=24,
            cost_microusd=None,
        )
    assert failure.value.code == "agent_cancelled"

    proposal = AgentMemoryProposal(
        namespace="preferences",
        key="blocked_run_preference",
        value="Use concise project updates.",
        language="en",
        source_ids=["goal"],
    )
    with pytest.raises(CoworkerError) as failure:
        container.memory.propose_from_agent(owner, run["id"], proposal, source_map)
    assert failure.value.code == "agent_cancelled"
    with container.repository.sessions() as db:
        assert db.scalars(select(AgentDecision).where(
            AgentDecision.run_id == run["id"],
        )).all() == []
        assert db.scalars(select(MemoryProposal).where(
            MemoryProposal.run_id == run["id"],
        )).all() == []

