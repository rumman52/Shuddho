"""Editable Office files from validated data; models never supply executable code."""
from __future__ import annotations

import csv
import html
import io
import json
import math
import unicodedata
import zipfile
from xml.etree import ElementTree as ET

from .calculations import aggregate_formula, row_formula, table_values
from .exports import RTL_LANGUAGES, font_for

INK, ACCENT, MUTED = "172C36", "4D46B8", "586C77"
FORMATS = {"text": "@", "number": "#,##0.00;-#,##0.00;0.00", "integer": "#,##0;-#,##0;0", "percent": "0.0%;-0.0%;0.0%"}


def text_height(text, width, font_size):
    """Conservative row height in points, including wide scripts and line breaks."""
    lines = 0
    for line in text.split("\n"):
        units = sum(2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1 for char in line)
        lines += max(1, math.ceil(units * font_size * .62 / (width * 5.25)))
    return max(font_size * 1.5, lines * font_size * 1.45 + 8)


def type_empty_formula_results(body):
    """Mark a known empty-string cache distinctly from a missing numeric cache.

    XlsxWriter emits <v/> without a cell type for value="". OOXML t="str"
    preserves its actual type, so a later safe import need not infer/evaluate it.
    Only our own newly generated, bounded package is rewritten here.
    """
    output = io.BytesIO()
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with zipfile.ZipFile(io.BytesIO(body)) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename.startswith("xl/worksheets/sheet") and entry.filename.endswith(".xml"):
                root = ET.fromstring(data)
                for cell in root.iter(namespace + "c"):
                    cache = cell.find(namespace + "v")
                    if cell.find(namespace + "f") is not None and cache is not None and cache.text is None:
                        cell.set("t", "str")
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(entry, data)
    return output.getvalue()


def display(value, kind):
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if kind == "percent":
        return f"{value * 100:,.1f}%"
    return f"{value:,.0f}" if kind == "integer" else f"{value:,.2f}"


def pdf_file(title, body, language, *, landscape=False):
    from pypdf import PdfReader
    from weasyprint import HTML
    direction = "rtl" if language.split("-")[0].lower() in RTL_LANGUAGES else "ltr"
    markup = f'''<!doctype html><html lang="{html.escape(language, quote=True)}" dir="{direction}"><meta charset="utf-8"><style>
      @page {{ size: A4 {'landscape' if landscape else ''}; margin: 16mm; @bottom-center {{ content: counter(page); font-size: 9pt; }} }}
      body {{ font-family: '{font_for(language)}', 'Noto Sans', 'DejaVu Sans', sans-serif; color: #{INK}; font-size: 11pt; line-height: 1.4; overflow-wrap: anywhere; }}
      h1 {{ font-size: 24pt; line-height: 1.2; }} h2 {{ font-size: 17pt; break-after: avoid; }}
      p, li {{ white-space: pre-wrap; }} section {{ break-before: page; }} section:first-of-type {{ break-before: auto; }}
      table {{ border-collapse: collapse; width: 100%; table-layout: fixed; font-size: 9pt; }}
      th, td {{ text-align: start; padding: 6px; border-bottom: 1px solid #ccd4d9; vertical-align: top; overflow-wrap: anywhere; }}
      th {{ color: white; background: #{INK}; }} thead {{ display: table-header-group; }} tr {{ break-inside: avoid; }}
      tfoot td {{ font-weight: bold; border-top: 2px solid #{INK}; }} .summary {{ color: #{MUTED}; }}
    </style><h1>{html.escape(title)}</h1>{body}</html>'''

    def no_resources(*args, **kwargs):
        raise ValueError("External resources are disabled")

    result = HTML(string=markup, url_fetcher=no_resources).write_pdf()
    if not PdfReader(io.BytesIO(result)).pages:
        raise ValueError("PDF export is empty")
    return result


def presentation_files(draft):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_TICK_LABEL_POSITION, XL_LEGEND_POSITION
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    deck.core_properties.title = draft.title
    deck.core_properties.language = draft.output_language
    deck.core_properties.author = "Shuddho"
    rtl = draft.output_language.split("-")[0].lower() in RTL_LANGUAGES

    def textbox(slide, text, x, y, width, height, size, *, bold=False, color=INK, bullet=False, direction=None):
        right_to_left = rtl if direction is None else direction
        frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height)).text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        for index, line in enumerate(text if isinstance(text, list) else [text]):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name = font_for(draft.output_language)
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
            paragraph.space_after = Pt(16 if bullet else 8)
            paragraph.line_spacing = 1.1
            paragraph.alignment = PP_ALIGN.RIGHT if right_to_left else PP_ALIGN.LEFT
            properties = paragraph._p.get_or_add_pPr()
            properties.set("rtl", "1" if right_to_left else "0")
            if bullet:
                properties.set("marL", str(Pt(24)))
                properties.set("indent", str(-Pt(18)))
                marker = OxmlElement("a:buChar")
                marker.set("char", "•")
                properties.append(marker)
            # Complex and East Asian scripts need explicit typefaces in Office.
            for run in paragraph.runs:
                props = run._r.get_or_add_rPr()
                props.set("lang", draft.output_language)
                for name in ("a:latin", "a:ea", "a:cs"):
                    element = OxmlElement(name)
                    element.set("typeface", font_for(draft.output_language))
                    props.append(element)
        return frame

    handout, transcript = [], [draft.title]
    for index, item in enumerate(draft.slides):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
        textbox(slide, f"{index + 1:02d} / {len(draft.slides):02d}", .75, .38, 2, .35, 12, color=ACCENT, direction=False)
        if item.layout == "cover":
            textbox(slide, item.title, .8, 1.4, 11.7, 3.0, 44, bold=True)
            textbox(slide, item.bullets, .8, 4.7, 11.7, 1.65, 24, color=MUTED)
        else:
            textbox(slide, item.title, .8, .93, 11.7, 1.72, 32, bold=True)
            if item.layout == "content":
                textbox(slide, item.bullets, .8, 2.85, 11.7, 3.92, 22, bullet=True)
            else:
                values = CategoryChartData()
                values.categories = item.chart.categories
                for series in item.chart.series:
                    values.add_series(series.label, series.values)
                chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED if item.chart.kind == "bar" else XL_CHART_TYPE.LINE_MARKERS,
                    Inches(.8), Inches(2.75), Inches(11.7), Inches(2.75), values).chart
                chart.font.name = font_for(draft.output_language)
                chart.font.size = Pt(14)
                chart.has_legend = True
                chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                chart.legend.include_in_layout = False
                chart.value_axis.has_title = bool(item.chart.unit)
                if item.chart.unit:
                    chart.value_axis.axis_title.text_frame.text = item.chart.unit
                chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
                textbox(slide, item.bullets, .8, 5.78, 11.7, 1.05, 18, color=MUTED)
        slide.notes_slide.notes_text_frame.text = item.speaker_notes
        body = f"<section><h2>{index + 1}. {html.escape(item.title)}</h2>"
        body += "<ul>" + "".join("<li>" + html.escape(line) + "</li>" for line in item.bullets) + "</ul>"
        chart_text = []
        if item.chart:
            chart_text = [item.chart.unit, "\t".join(["", *[series.label for series in item.chart.series]])]
            body += "<p>" + html.escape(item.chart.unit) + "</p><table><thead><tr><th></th>"
            body += "".join("<th>" + html.escape(series.label) + "</th>" for series in item.chart.series) + "</tr></thead><tbody>"
            for row, category in enumerate(item.chart.categories):
                numbers = [str(series.values[row]) for series in item.chart.series]
                chart_text.append("\t".join([category, *numbers]))
                body += "<tr>" + "".join("<td>" + html.escape(value) + "</td>" for value in [category, *numbers]) + "</tr>"
            body += "</tbody></table>"
        body += f"<p>{html.escape(item.speaker_notes)}</p></section>"
        handout.append(body)
        transcript.extend([f"{index + 1}. {item.title}", *item.bullets, *chart_text, item.speaker_notes])
    buffer = io.BytesIO()
    deck.save(buffer)
    body = buffer.getvalue()
    if len(Presentation(io.BytesIO(body)).slides) != len(draft.slides):
        raise ValueError("Slide count changed during export")
    return [("presentation.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation", body),
            ("presentation.pdf", "application/pdf", pdf_file(draft.title, "".join(handout), draft.output_language)),
            ("presentation.txt", "text/plain; charset=utf-8", "\n\n".join(transcript).encode("utf-8"))]


def spreadsheet_files(draft):
    import xlsxwriter
    from openpyxl import load_workbook
    from xlsxwriter.utility import xl_rowcol_to_cell

    table = table_values(draft)
    columns = draft.columns + draft.calculations
    output = io.BytesIO()
    book = xlsxwriter.Workbook(output, {"in_memory": True, "strings_to_formulas": False, "strings_to_urls": False})
    book.set_properties({"title": draft.title, "author": "Shuddho"})
    book.set_calc_mode("auto")
    sheet = book.add_worksheet(draft.sheet_name)
    rtl = draft.output_language.split("-")[0].lower() in RTL_LANGUAGES
    if rtl:
        sheet.right_to_left()
    base = {"font_name": font_for(draft.output_language), "font_size": 11, "valign": "top"}
    title = book.add_format(base | {"font_size": 20, "bold": True, "font_color": "#" + INK, "text_wrap": True})
    summary = book.add_format(base | {"text_wrap": True, "font_color": "#" + MUTED})
    header = book.add_format(base | {"bold": True, "bg_color": "#" + INK, "font_color": "white", "text_wrap": True})
    # Inputs are blue; formulas are ink. Typed strings can never become formulas.
    formats = [book.add_format(base | {"num_format": FORMATS[c.format], "text_wrap": c.format == "text",
               "align": "right" if rtl or c.format != "text" else "left", "indent": 1,
               "font_color": "#215C9E" if i < len(draft.columns) else "#" + INK}) for i, c in enumerate(columns)]
    totals = [book.add_format(base | {"num_format": FORMATS[c.format], "bold": True, "top": 2, "font_color": "#" + INK}) for c in columns]
    end = max(1, len(columns) - 1)
    widths = [28 if column.format == "text" else max(22, max(len(display(row[i], column.format)) for row in table["rows"]) * 1.2 + 2)
              for i, column in enumerate(columns)]
    merged_width = sum(widths) + (28 if len(columns) == 1 else 0)
    if merged_width < 80:
        widths[0] += 80 - merged_width
        merged_width = 80
    if len(columns) == 1:
        sheet.set_column(1, 1, 28)
    sheet.merge_range(0, 0, 0, end, draft.title, title)
    sheet.set_row(0, max(56, text_height(draft.title, merged_width, 20)))
    # A single wrapped row moves together at a page boundary. Merging vertically
    # across rows can cut a line in half when a long summary crosses a page.
    sheet.merge_range(1, 0, 1, end, draft.summary, summary)
    sheet.set_row(1, max(24, text_height(draft.summary, merged_width, 11)))
    sheet.set_row(2, 12)
    sheet.set_row(3, 12)
    header_row, first = 4, 5  # Excel header row 5; data starts on row 6.
    sheet.set_row(header_row, max(48, *(text_height(column.label, widths[i], 11) for i, column in enumerate(columns))))
    for index, column in enumerate(columns):
        sheet.write_string(header_row, index, column.label, header)
        sheet.set_column(index, index, widths[index])
    for offset, row in enumerate(table["rows"]):
        number = first + offset
        addresses = {column.id: xl_rowcol_to_cell(number, i) for i, column in enumerate(columns)}
        sheet.set_row(number, max(44, *(text_height(value, widths[i], 11) if isinstance(value, str) else 0 for i, value in enumerate(row))))
        for index, (column, value) in enumerate(zip(columns, row)):
            if index >= len(draft.columns):
                sheet.write_formula(number, index, row_formula(column, addresses), formats[index], "" if value is None else value)
            elif value is None:
                sheet.write_blank(number, index, None, formats[index])
            elif isinstance(value, str):
                sheet.write_string(number, index, value, formats[index])
            else:
                sheet.write_number(number, index, value, formats[index])
    last = first + len(draft.rows) - 1
    total_row = last + 2
    if draft.summary_label:
        sheet.merge_range(total_row, 0, total_row, end, draft.summary_label, summary)
        sheet.set_row(total_row, text_height(draft.summary_label, merged_width, 11))
        for index, column in enumerate(columns):
            if column.aggregate != "none":
                cached = table["totals"][index]
                sheet.write_formula(total_row + 1, index, aggregate_formula(column, xl_rowcol_to_cell(first, index),
                    xl_rowcol_to_cell(last, index), len(draft.rows)), totals[index], "" if cached is None else cached)
    sheet.freeze_panes(first, 0)
    sheet.autofilter(header_row, 0, last, len(columns) - 1)
    sheet.hide_gridlines(2)
    sheet.set_landscape()
    sheet.set_paper(9)
    sheet.set_margins(.5, .5, .5, .5)
    sheet.fit_to_pages(1, 0)
    sheet.repeat_rows(header_row)
    chart_end = total_row + 2
    if draft.chart:
        # fit_to_pages overrides manual breaks. Set an explicit width scale so
        # a chart starts intact on the next printed page rather than splitting.
        print_width = sum(width * 5.25 + 4 for width in widths)
        if print_width > 740:
            sheet.set_paper(8)  # A3 landscape for wider tables.
        sheet.set_print_scale(min(100, int((1100 if print_width > 740 else 740) / print_width * 100)))
        sheet.set_h_pagebreaks([total_row + 3])
        indexes = {column.id: index for index, column in enumerate(columns)}
        chart = book.add_chart({"type": "column" if draft.chart.kind == "bar" else "line"})
        for key in draft.chart.series:
            index = indexes[key]
            chart.add_series({"name": [draft.sheet_name, header_row, index],
                              "categories": [draft.sheet_name, first, indexes[draft.chart.category], last, indexes[draft.chart.category]],
                              "values": [draft.sheet_name, first, index, last, index]})
        chart.set_title({"name": draft.chart.title, "name_font": {"name": font_for(draft.output_language), "size": 16}})
        chart_font = {"name": font_for(draft.output_language), "size": 11}
        chart.set_legend({"position": "bottom", "font": chart_font})
        chart.set_x_axis({"num_font": chart_font})
        chart.set_y_axis({"num_font": chart_font})
        chart.show_blanks_as("gap")
        chart.set_style(10)
        sheet.insert_chart(total_row + 3, 0, chart, {"x_scale": 1.35, "y_scale": 1.15})
        chart_end += 21
    sheet.print_area(0, 0, chart_end, end)
    book.close()
    body = type_empty_formula_results(output.getvalue())
    check = load_workbook(io.BytesIO(body), read_only=True, data_only=False)
    try:
        if check.sheetnames != [draft.sheet_name]:
            raise ValueError("Workbook export could not be reopened")
    finally:
        check.close()
    csv_output = io.StringIO(newline="")
    writer = csv.writer(csv_output)

    def csv_value(value):
        # CSV has no cell types. Neutralize text that spreadsheet apps execute.
        if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
            return "'" + value
        return "" if value is None else value

    writer.writerow([csv_value(column.label) for column in columns])
    writer.writerows([csv_value(value) for value in row] for row in table["rows"])
    markup = f'<p class="summary">{html.escape(draft.summary)}</p><table><thead><tr>'
    markup += "".join("<th>" + html.escape(column.label) + "</th>" for column in columns) + "</tr></thead><tbody>"
    for row in table["rows"]:
        markup += "<tr>" + "".join("<td>" + html.escape(display(value, column.format)) + "</td>" for column, value in zip(columns, row)) + "</tr>"
    markup += "</tbody></table>"
    if draft.summary_label:
        markup += f"<p>{html.escape(draft.summary_label)}</p><table><tr>" + "".join("<td>" + html.escape(display(value, column.format)) + "</td>" for column, value in zip(columns, table["totals"])) + "</tr></table>"
    return [("spreadsheet.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", body),
            ("spreadsheet.pdf", "application/pdf", pdf_file(draft.title, markup, draft.output_language, landscape=True)),
            ("spreadsheet.csv", "text/csv; charset=utf-8", csv_output.getvalue().encode("utf-8-sig"))]


def render_office_artifacts(skill_id, draft, sources):
    files = presentation_files(draft) if skill_id == "presentation" else spreadsheet_files(draft)
    files.append(("source-manifest.json", "application/json", json.dumps({
        "skill_id": skill_id, "sources": sources, "missing_information": draft.missing_information,
    }, ensure_ascii=False, indent=2).encode("utf-8")))
    return files
