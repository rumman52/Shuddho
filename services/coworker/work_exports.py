"""Native work-service layouts. Plain data enters, verified files leave."""
from __future__ import annotations

import html
import io
import json
from dataclasses import dataclass

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader

from .schemas import DocumentPackage, EmailPackage, MeetingPackage, PlanPackage, ResearchPackage, SocialPackage, parse_draft
from .skills import SKILLS


@dataclass(frozen=True)
class Block:
    text: str
    kind: str = "paragraph"
    url: str | None = None


def content_blocks(draft, sources=None):
    blocks = []
    if isinstance(draft, ResearchPackage):
        from .research import public_source_url
        title = draft.title
        by_id = {source["id"]: source for source in sources or [] if source.get("kind") == "web"}
        cited = []
        for finding in draft.findings:
            blocks.extend([Block(finding.heading, "heading"), Block(finding.text)])
            for citation in finding.citations:
                if citation.source_id not in by_id:
                    raise ValueError("Missing research source")
                if citation.source_id not in cited:
                    cited.append(citation.source_id)
                number = cited.index(citation.source_id) + 1
                blocks.append(Block(f'{draft.labels.evidence} [{number}]: “{citation.quote}”', "meta"))
        if draft.missing_information:
            blocks.append(Block(draft.labels.gaps, "heading"))
            blocks.extend(Block(text, "bullet") for text in draft.missing_information)
        if cited:
            blocks.append(Block(draft.labels.sources, "heading"))
        for index, source_id in enumerate(cited, 1):
            source = by_id[source_id]
            url = public_source_url(source["url"])
            blocks.append(Block(f'[{index}] {source["label"]}', "reference", url))
            blocks.append(Block(url, "url", url))
            blocks.append(Block(f'{draft.labels.retrieved}: {source["retrieved_at"][:10]} · '
                                f'{draft.labels.source_date}: {source["source_date"][:10] if source.get("source_date") else draft.labels.undated}', "meta"))
    elif isinstance(draft, DocumentPackage):
        title = draft.document.title
        if draft.document.summary:
            blocks.append(Block(draft.document.summary))
        for section in draft.document.sections:
            if section.heading:
                blocks.append(Block(section.heading, "heading"))
            blocks.extend(Block(text) for text in section.paragraphs)
            blocks.extend(Block(text, "bullet") for text in section.bullets)
    elif isinstance(draft, EmailPackage):
        title = draft.email.subject
        blocks.append(Block(draft.email.body))
    elif isinstance(draft, SocialPackage):
        title = draft.title
        for post in draft.posts:
            blocks.append(Block(post.label, "heading"))
            blocks.append(Block({"facebook": "Facebook", "linkedin": "LinkedIn", "other": ""}[post.platform], "meta"))
            blocks.append(Block(post.text))
            if post.suggested_timing:
                blocks.append(Block(post.suggested_timing, "meta"))
    elif isinstance(draft, MeetingPackage):
        title = draft.title
        blocks.extend([Block(draft.summary), Block(draft.labels.notes, "heading")])
        blocks.extend(Block(note.text, "bullet") for note in draft.notes)
        if draft.decisions:
            blocks.append(Block(draft.labels.decisions, "heading"))
            blocks.extend(Block(item.text, "bullet") for item in draft.decisions)
        if draft.actions:
            blocks.append(Block(draft.labels.actions, "heading"))
        for action in draft.actions:
            blocks.append(Block(action.text, "check"))
            metadata = [getattr(draft.labels, action.basis)]
            if action.owner:
                metadata.append(draft.labels.owner + ": " + action.owner)
            if action.deadline:
                metadata.append(draft.labels.deadline + ": " + action.deadline)
            blocks.append(Block(" · ".join(metadata), "meta"))
    elif isinstance(draft, PlanPackage):
        title = draft.title
        blocks.append(Block(draft.overview))
        for item in draft.items:
            if item.when:
                blocks.append(Block(item.when, "heading"))
            blocks.append(Block(item.task, "check"))
            blocks.append(Block(getattr(draft.labels, item.basis) + " · " + getattr(draft.labels, item.priority), "meta"))
    else:
        raise ValueError("Unknown work-service draft")
    return title, [block for block in blocks if block.text]


def render_work_artifacts(skill_id, draft, sources):
    # Validate the requested service again inside the renderer process. Filenames,
    # markup and object destinations are never accepted from model output.
    draft = parse_draft(skill_id, draft.model_dump())
    title, blocks = content_blocks(draft, sources)
    docx, pdf = render_document(title, blocks, draft.output_language)
    stem = SKILLS[skill_id].filename
    text = title + "\n\n" + "\n\n".join(
        ("• " if block.kind == "bullet" else "☐ " if block.kind == "check" else "") + block.text for block in blocks)
    manifest = {"skill_id": skill_id, "output_language": draft.output_language,
                "sources": sources, "missing_information": draft.missing_information}
    return [
        (stem + ".docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", docx),
        (stem + ".pdf", "application/pdf", pdf),
        (stem + ".txt", "text/plain; charset=utf-8", text.encode("utf-8")),
        ("source-manifest.json", "application/json", json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8")),
    ]


def render_document(title: str, blocks: list[Block], language: str):
    from .exports import RTL_LANGUAGES, font_for
    from weasyprint import HTML

    rtl = language.split("-")[0].lower() in RTL_LANGUAGES
    font = font_for(language)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.85)
    for name in ("Normal", "Title", "Heading 1", "List Bullet"):
        style = doc.styles[name]
        style.font.name = font
        style.font.color.rgb = RGBColor.from_string("28353E")
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:" + key), font)
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.35
    doc.styles["Normal"].paragraph_format.space_after = Pt(9)

    def paragraph(text, style=None):
        value = doc.add_paragraph(text, style)
        if rtl:
            bidi = OxmlElement("w:bidi")
            bidi.set(qn("w:val"), "1")
            value._p.get_or_add_pPr().append(bidi)
        for run in value.runs:
            lang = OxmlElement("w:lang")
            lang.set(qn("w:val"), language)
            lang.set(qn("w:bidi"), language)
            run._r.get_or_add_rPr().append(lang)
            if rtl:
                run.font.rtl = True
        return value

    paragraph(title, "Title")
    markup = [f"<h1>{html.escape(title)}</h1>"]
    for block in blocks:
        text = ("☐ " if block.kind == "check" else "") + block.text
        value = paragraph(text, "Heading 1" if block.kind == "heading" else "List Bullet" if block.kind == "bullet" else None)
        if block.kind == "meta":
            for run in value.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor.from_string("62717B")
        if block.url:
            # Only server-owned validated source URLs become native hyperlinks.
            from .research import public_source_url
            url = public_source_url(block.url)
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), value.part.relate_to(url, RT.HYPERLINK, is_external=True))
            for run in list(value.runs):
                run.font.color.rgb = RGBColor.from_string("423570")
                run.font.underline = True
                if block.kind == "url":
                    run.font.rtl = False
                hyperlink.append(run._r)
            value._p.append(hyperlink)
            if block.kind == "url":
                bidi = value._p.get_or_add_pPr().find(qn("w:bidi"))
                if bidi is not None:
                    bidi.set(qn("w:val"), "0")
        tag = "h2" if block.kind == "heading" else "p"
        escaped = html.escape(text)
        if block.kind == "bullet":
            escaped = "• " + escaped
        if block.url:
            escaped = f'<a href="{html.escape(url, quote=True)}">{escaped}</a>'
        markup.append(f"<{tag} class='{block.kind}'>{escaped}</{tag}>")
    buffer = io.BytesIO()
    doc.save(buffer)
    docx = buffer.getvalue()
    if not Document(io.BytesIO(docx)).paragraphs:
        raise ValueError("Empty DOCX export")
    direction = "rtl" if rtl else "ltr"
    page = f"""<!doctype html><html lang="{html.escape(language, quote=True)}" dir="{direction}"><head><meta charset="utf-8"><style>
      @page {{ size: A4; margin: 20mm 22mm; @bottom-center {{ content: counter(page); font-size: 9pt; color: #62717b; }} }}
      body {{ font-family: '{font}', 'Noto Sans', 'DejaVu Sans', sans-serif; font-size: 11pt; line-height: 1.6; color: #28353e; overflow-wrap: anywhere; }}
      h1 {{ font-size: 25pt; line-height: 1.3; margin: 0 0 7mm; color: #423570; }}
      h2 {{ font-size: 13pt; margin: 6mm 0 2mm; break-after: avoid; }}
      p {{ white-space: pre-wrap; orphans: 3; widows: 3; margin: 0 0 3mm; }}
      .meta {{ font-size: 9pt; color: #62717b; }} .bullet, .check {{ padding-inline-start: 3mm; }}
      a {{ color: #423570; }} .url {{ font-size: 9pt; direction: ltr; text-align: left; word-break: break-all; }}
    </style></head><body>{''.join(markup)}</body></html>"""

    def no_external_resources(*_args, **_kwargs):
        raise ValueError("External resources are disabled")

    pdf = HTML(string=page, url_fetcher=no_external_resources).write_pdf()
    if not pdf.startswith(b"%PDF-") or not PdfReader(io.BytesIO(pdf)).pages:
        raise ValueError("Invalid PDF export")
    return docx, pdf
