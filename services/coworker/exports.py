from __future__ import annotations

import html
import io
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader

from .schemas import AnyDraft, DraftPackage
from .skills import SkillId, artifact_filenames
from .errors import CoworkerError

RTL_LANGUAGES = {"ar", "fa", "he", "ur", "ps", "dv"}


def render_in_subprocess(draft: AnyDraft, sources: list[dict], *, skill_id: SkillId = "report_email"):
    with tempfile.TemporaryDirectory(prefix="shuddho-export-") as directory:
        folder = Path(directory)
        (folder / "input.json").write_text(json.dumps({"draft": draft.model_dump(), "sources": sources, "skill_id": skill_id}, ensure_ascii=False), encoding="utf-8")
        try:
            result = subprocess.run(
                [sys.executable, "-m", "services.coworker.render_worker", directory],
                cwd=Path(__file__).resolve().parents[2], env={"PATH": os.defpath, "LANG": "C.UTF-8"},
                stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=60,
            )
            manifest_path = folder / "manifest.json"
            if result.returncode or not manifest_path.is_file() or manifest_path.stat().st_size > 2000:
                raise ValueError()
            manifest = json.loads(manifest_path.read_text())
            expected = artifact_filenames(skill_id)
            if {row[0] for row in manifest} != expected or len(manifest) != 4:
                raise ValueError()
            outputs = []
            for filename, content_type in manifest:
                path = folder / filename
                if not 0 < path.stat().st_size <= 16 * 1024 * 1024:
                    raise ValueError()
                outputs.append((filename, content_type, path.read_bytes()))
            return outputs
        except (subprocess.TimeoutExpired, ValueError, OSError, TypeError):
            raise CoworkerError("export_failed", "The file export could not finish. Try a shorter draft.", 422) from None


def font_for(language):
    root = language.split("-")[0].lower()
    return {"bn": "Noto Sans Bengali", "hi": "Noto Sans Devanagari", "ar": "Noto Sans Arabic",
            "fa": "Noto Sans Arabic", "ur": "Noto Sans Arabic", "he": "Noto Sans Hebrew",
            "zh": "Noto Sans CJK SC", "ja": "Noto Sans CJK JP", "ko": "Noto Sans CJK KR"}.get(root, "Noto Sans")


def render_artifacts(draft: DraftPackage, sources: list[dict]) -> list[tuple[str, str, bytes]]:
    language = draft.output_language
    rtl = language.split("-")[0].lower() in RTL_LANGUAGES
    report = draft.report
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = document.styles[name]
        style.font.name = font_for(language)
        style.font.color.rgb = RGBColor.from_string("173A39")
        for key in ["ascii", "hAnsi", "eastAsia", "cs"]:
            style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:" + key), font_for(language))
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.3

    def paragraph(text, style=None):
        value = document.add_paragraph(text, style)
        props = value._p.get_or_add_pPr()
        if rtl:
            element = OxmlElement("w:bidi")
            element.set(qn("w:val"), "1")
            props.append(element)
        for run in value.runs:
            lang = OxmlElement("w:lang")
            lang.set(qn("w:val"), language)
            lang.set(qn("w:bidi"), language)
            run._r.get_or_add_rPr().append(lang)
            if rtl:
                run.font.rtl = True
        return value

    paragraph(report.title, "Title")
    paragraph(report.summary)
    parts = [f"<h1>{html.escape(report.title)}</h1><p class='summary'>{html.escape(report.summary)}</p>"]
    for item in report.sections:
        paragraph(item.heading, "Heading 1")
        parts.append(f"<h2>{html.escape(item.heading)}</h2>")
        for text in item.paragraphs:
            paragraph(text)
            parts.append(f"<p>{html.escape(text)}</p>")
        references = "[" + ", ".join(item.source_ids) + "]"
        paragraph(references)
        parts.append(f"<p class='reference'>{html.escape(references)}</p>")
    for source in sources:
        reference = f'[{source["id"]}] {source["label"]}'
        paragraph(reference)
        parts.append(f"<p class='reference'>{html.escape(reference)}</p>")
    if draft.missing_information:
        for item in draft.missing_information:
            paragraph("[ ] " + item)
            parts.append(f"<p class='missing'>□ {html.escape(item)}</p>")
    buffer = io.BytesIO()
    document.save(buffer)
    docx = buffer.getvalue()
    # Re-open the generated package before exposing it as a successful export.
    if not Document(io.BytesIO(docx)).paragraphs:
        raise ValueError("DOCX export is empty")
    direction = "rtl" if rtl else "ltr"
    markup = f"""<!doctype html><html lang="{html.escape(language, quote=True)}" dir="{direction}"><head><meta charset="utf-8"><style>
      @page {{ size: A4; margin: 20mm 22mm; @bottom-center {{ content: counter(page); font-size: 9pt; color: #667; }} }}
      body {{ font-family: '{font_for(language)}', 'Noto Sans', 'DejaVu Sans', sans-serif; font-size: 11pt; line-height: 1.55; color: #203636; overflow-wrap: anywhere; }}
      h1 {{ font-size: 26pt; line-height: 1.22; color: #143f3c; margin: 0 0 7mm; }}
      h2 {{ font-size: 14pt; margin: 7mm 0 2mm; break-after: avoid; }}
      p {{ orphans: 3; widows: 3; margin: 0 0 3mm; white-space: pre-wrap; }}
      .summary {{ font-size: 12pt; color: #455c5b; margin-bottom: 7mm; }}
      .reference {{ font-size: 8.5pt; color: #667; }} .missing {{ color: #72501c; }}
    </style></head><body>{''.join(parts)}</body></html>"""
    from weasyprint import HTML

    def no_external_resources(*_args, **_kwargs):
        raise ValueError("External resources are disabled for document export")

    pdf = HTML(string=markup, url_fetcher=no_external_resources).write_pdf()
    if not pdf.startswith(b"%PDF-") or not PdfReader(io.BytesIO(pdf)).pages:
        raise ValueError("PDF export is invalid")
    email = draft.email.subject + "\n\n" + draft.email.body
    return [
        ("report.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", docx),
        ("report.pdf", "application/pdf", pdf),
        ("email-draft.txt", "text/plain; charset=utf-8", email.encode("utf-8")),
        ("source-manifest.json", "application/json", json.dumps({"sources": sources}, ensure_ascii=False, indent=2).encode("utf-8")),
    ]
