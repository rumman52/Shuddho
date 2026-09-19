"""Bounded Office/CSV snapshots. Never execute formulas, links, or embedded code."""
from __future__ import annotations

import csv
import io
import json
import posixpath
import re
import zipfile
from datetime import date, datetime, time
from xml.etree import ElementTree as ET

from .errors import CoworkerError

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def invalid(message):
    return CoworkerError("document_invalid", message, 415)


def safe_archive(data, kind):
    archive = zipfile.ZipFile(io.BytesIO(data))
    members = archive.infolist()
    names = {item.filename for item in members}
    if len(members) > 2000 or sum(item.file_size for item in members) > 32 * 1024 * 1024:
        archive.close()
        raise CoworkerError("document_expansion_limit", "This file is too large to unpack safely.", 413)
    try:
        if len(names) != len(members):
            raise invalid("This file contains duplicate package entries.")
        for item in members:
            name = item.filename.lower()
            if (item.flag_bits & 1 or name.startswith("/") or "\\" in name or ".." in name.split("/")
                    or any(part in name for part in ("vbaproject", "/activex/", "externallinks/", "connections.xml"))
                    or (name.endswith(".bin") and not re.fullmatch(r"ppt/printersettings/printersettings\d+\.bin", name))
                    or ("/embeddings/" in name and not (kind == "pptx" and name.endswith(".xlsx")))):
                raise invalid("Upload a standard file without macros, external data connections, encryption, or embedded objects.")
            if name.endswith((".xml", ".rels")):
                root = xml(archive, item.filename)
                if kind == "xlsx" and name.startswith("xl/worksheets/") and name.endswith(".xml"):
                    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                    for cell in root.iter(ns + "c"):
                        address = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", cell.get("r", ""))
                        if not address or len(address[1]) > 1 or address[1] > "T" or int(address[2]) > 201:
                            raise CoworkerError("table_limit", "Choose a workbook with at most 200 data rows, one header row, and 20 columns per sheet.", 413)
        if "[Content_Types].xml" not in names:
            raise invalid("This is not a standard Office file.")
        return archive
    except BaseException:
        archive.close()
        raise


def xml(archive, path):
    content = archive.read(path)
    if len(content) > 8 * 1024 * 1024 or b"<!DOCTYPE" in content.upper() or b"<!ENTITY" in content.upper():
        raise invalid("This file contains unsupported XML content.")
    return ET.fromstring(content)


def relationships(archive, path):
    relpath = posixpath.join(posixpath.dirname(path), "_rels", posixpath.basename(path) + ".rels")
    if relpath not in archive.namelist():
        return {}
    result = {}
    for rel in xml(archive, relpath):
        if rel.get("TargetMode") == "External":
            continue  # No network access; hyperlinks are never followed.
        target = rel.get("Target", "")
        target = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join(posixpath.dirname(path), target))
        if target.startswith("../") or "\\" in target or target not in archive.namelist():
            raise invalid("This file contains an invalid relationship.")
        result[rel.get("Id")] = (target, rel.get("Type", "").rsplit("/", 1)[-1])
    return result


def pptx_snapshot(archive):
    main = "ppt/presentation.xml"
    rels = relationships(archive, main)
    slides = list(xml(archive, main).iter(P + "sldId"))
    if not 1 <= len(slides) <= 20:
        raise CoworkerError("slide_limit", "Upload a presentation with 1–20 slides.", 413)
    result = []
    for index, reference in enumerate(slides):
        path, kind = rels[reference.get(R + "id")]
        if kind != "slide":
            raise invalid("This presentation has an invalid slide reference.")
        root = xml(archive, path)
        slide_rels = relationships(archive, path)
        text = ["".join(node.text or "" for node in paragraph.iter(A + "t")) for paragraph in root.iter(A + "p")]
        charts = []
        for reference in root.iter(C + "chart"):
            chart_path, chart_kind = slide_rels[reference.get(R + "id")]
            if chart_kind != "chart":
                raise invalid("This presentation has an invalid chart reference.")
            series = []
            for item in xml(archive, chart_path).iter(C + "ser"):
                # Extract cached chart values only. Never open the embedded workbook.
                points = {}
                for key in ("tx", "cat", "val", "xVal", "yVal", "bubbleSize"):
                    element = item.find(C + key)
                    if element is None:
                        continue
                    if element.find(".//" + C + "f") is not None and element.find(".//" + C + "v") is None:
                        raise invalid("This chart has no saved data. Re-save the presentation with chart data before uploading.")
                    points[key] = [node.text or "" for node in element.iter(C + "v")]
                series.append(points)
            charts.append(series)
        notes = []
        for target, relation in slide_rels.values():
            if relation == "notesSlide":
                for shape in xml(archive, target).iter(P + "sp"):
                    placeholder = shape.find(".//" + P + "ph")
                    if placeholder is not None and placeholder.get("type") in {"sldNum", "hdr", "ftr", "dt"}:
                        continue
                    notes.extend("".join(node.text or "" for node in paragraph.iter(A + "t")) for paragraph in shape.iter(A + "p"))
        result.append({"slide": index + 1, "text": text, "speaker_notes": notes, "cached_chart_data": charts})
    return "Presentation snapshot: slide text, tables, speaker notes, and saved chart data only. Images, diagrams, and layout are not interpreted.\n" + json.dumps(result, ensure_ascii=False)


def xlsx_snapshot(data):
    from openpyxl import load_workbook
    values = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        if not 1 <= len(values.worksheets) <= 4:
            raise CoworkerError("sheet_limit", "Upload a workbook with 1–4 worksheets.", 413)
        output = []
        for sheet, source in zip(values.worksheets, formulas.worksheets):
            # Ignore untrusted dimensions and inspect actual bounded rows/cells.
            sheet.reset_dimensions()
            source.reset_dimensions()
            rows = []
            for row_index, (row, original) in enumerate(zip(sheet.iter_rows(), source.iter_rows())):
                if row_index >= 201 or len(row) > 20:
                    raise CoworkerError("table_limit", "Choose a workbook with at most 200 data rows, one header row, and 20 columns per sheet.", 413)
                cells = []
                for cell, expression in zip(row, original):
                    if cell.data_type == "e":
                        raise invalid("This workbook contains a spreadsheet error. Correct it before uploading.")
                    if expression.data_type == "f" and cell.value is None and cell.data_type not in {"str", "s", "inlineStr"}:
                        raise invalid("Some formulas have no saved result. Recalculate and save the workbook, or upload a CSV of its values.")
                    value = cell.value
                    if isinstance(value, (date, datetime, time)):
                        value = value.isoformat()
                    cells.append({"value": value, "number_format": cell.number_format}
                                 if isinstance(value, (int, float)) and not isinstance(value, bool) and cell.number_format != "General" else value)
                rows.append(cells)
            output.append({"sheet": sheet.title, "rows": rows})
        return "Workbook snapshot: saved cell values only; formulas were not recalculated and may be stale. Charts, images, formatting, and hidden/visible distinctions are not interpreted.\n" + json.dumps(output, ensure_ascii=False, allow_nan=False)
    finally:
        values.close()
        formulas.close()


def extract_office_text(data: bytes, kind: str):
    if kind == "csv":
        try:
            reader = csv.reader(io.StringIO(data.decode("utf-8-sig")), strict=True)
            rows = []
            for row in reader:
                if len(rows) >= 201 or not 1 <= len(row) <= 20:
                    raise CoworkerError("table_limit", "Choose a CSV with at most 200 data rows, one header row, and 20 columns.", 413)
                if rows and len(row) != len(rows[0]):
                    raise invalid("Use the same number of columns in each CSV row.")
                rows.append(row)
            if not rows or not any(cell.strip() for row in rows for cell in row):
                raise invalid("This CSV has no readable cells.")
            return "CSV snapshot: all cells are supplied text, not executable formulas.\n" + json.dumps(rows, ensure_ascii=False)
        except UnicodeDecodeError:
            raise invalid("Save this CSV as UTF-8 and upload it again.") from None
        except csv.Error:
            raise invalid("This CSV could not be read. Check its quoting and separators.") from None
    with safe_archive(data, kind) as archive:
        return pptx_snapshot(archive) if kind == "pptx" else xlsx_snapshot(data)
